"""
Планирование профилактических работ (ТЗ §3 «поддержка планирования профилактических работ», §8):
рабочее место инженера ТО.

План зоны складывается из трёх частей:
- регламент — единицы реестра, у которых подходит или прошёл срок ТО (последнее ТО + интервал);
- состояние — рекомендации системы по прогнозам, Data Health и фактическому состоянию (recommendations.py);
- запланировано — открытые заявки на работы с датой.
Инженер ставит работу в план на выбранную дату: создаётся черновик заявки на эту дату, дальше он
идёт обычным путём (утверждение руководителем → бригада). Бригада при выполнении отмечает фактическое
состояние; ТО обновляет дату последнего обслуживания в реестре.
"""

from __future__ import annotations

from datetime import date, datetime, time, timedelta

from django.db import transaction
from django.utils import timezone

from apps.assets.models import Equipment, EquipmentCondition
from apps.incidents.models import IncidentEvent
from apps.topology.selectors import scope_queryset

from .models import EquipmentInspection, MaintenanceRecommendation, WorkOrder, WorkType
from .recommendations import EQUIPMENT_WORK
from .services import next_number

HORIZON_DAYS = 90
OPEN_ORDERS = (
    WorkOrder.Status.DRAFT,
    WorkOrder.Status.APPROVED,
    WorkOrder.Status.SUBMITTED,
    WorkOrder.Status.IN_PROGRESS,
)
# Работы, после которых у единицы начинается новый межрегламентный интервал
MAINTENANCE_WORK = {
    WorkType.CALIBRATION,
    WorkType.PUMP_SERVICE,
    WorkType.VENTILATION,
    WorkType.POWER_CHECK,
    WorkType.CLEANING,
    WorkType.SENSOR_REPLACEMENT,
    WorkType.INSPECTION,
    WorkType.SECURITY,
}
BAD = (EquipmentCondition.NEEDS_REPAIR, EquipmentCondition.FAULTY)


def work_for(eq: Equipment) -> str:
    if eq.kind == "sensor" and "метан" in eq.name.lower():
        return WorkType.CALIBRATION
    return EQUIPMENT_WORK.get(eq.kind, WorkType.INSPECTION)


def _equipment_row(eq: Equipment, today: date, planned: dict[int, WorkOrder]) -> dict:
    due = eq.next_maintenance_at
    order = planned.get(eq.pk)
    return {
        "id": eq.pk,
        "name": eq.name,
        "kind": eq.kind,
        "kind_display": eq.get_kind_display(),
        "inventory_number": eq.inventory_number,
        "node": eq.node_id,
        "node_name": eq.node.name,
        "picket": float(eq.picket) if eq.picket is not None else None,
        "last_maintenance_at": eq.last_maintenance_at,
        "interval_days": eq.maintenance_interval_days,
        "due": due,
        "overdue_days": (today - due).days if due and due < today else 0,
        "condition": eq.condition,
        "condition_display": eq.get_condition_display() if eq.condition else "",
        "condition_at": eq.condition_at,
        "source": eq.source,
        "work_type": work_for(eq),
        "planned": {"id": order.pk, "number": order.number, "status": order.status, "due_at": order.due_at}
        if order
        else None,
    }


def plan(user, horizon_days: int = HORIZON_DAYS, today: date | None = None) -> dict:
    today = today or timezone.localdate()
    until = today + timedelta(days=horizon_days)
    equipment = scope_queryset(Equipment.objects.filter(is_active=True), user, "node").select_related("node")
    orders = (
        scope_queryset(WorkOrder.objects.filter(status__in=OPEN_ORDERS), user, "node")
        .select_related("node", "equipment", "assignee")
        .order_by("due_at")
    )
    planned_by_eq = {o.equipment_id: o for o in orders if o.equipment_id}
    due_rows = []
    for eq in equipment.exclude(maintenance_interval_days=None):
        due = eq.next_maintenance_at
        if due is not None and due <= until:
            due_rows.append(_equipment_row(eq, today, planned_by_eq))
    due_rows.sort(key=lambda r: r["due"])
    bad_rows = [_equipment_row(eq, today, planned_by_eq) for eq in equipment.filter(condition__in=BAD)]
    recs = (
        scope_queryset(
            MaintenanceRecommendation.objects.filter(status=MaintenanceRecommendation.Status.NEW),
            user,
            "node",
        )
        .select_related("node", "equipment", "channel")
        .order_by("due_date")
    )
    weeks: dict[date, int] = {}
    for o in orders:
        week = timezone.localtime(o.due_at).date()
        week -= timedelta(days=week.weekday())
        weeks[week] = weeks.get(week, 0) + 1
    overdue = [r for r in due_rows if r["overdue_days"] > 0]
    return {
        "today": today,
        "horizon_days": horizon_days,
        "kpis": {
            "overdue": len(overdue),
            "due_30": sum(
                1 for r in due_rows if not r["overdue_days"] and r["due"] <= today + timedelta(days=30)
            ),
            "unplanned_overdue": sum(1 for r in overdue if not r["planned"]),
            "recommendations": recs.count(),
            "bad_condition": len(bad_rows),
            "planned": len(orders),
        },
        "due": due_rows[:300],
        "bad_condition": bad_rows[:100],
        "recommendations": [
            {
                "id": r.pk,
                "node_name": r.node.name,
                "equipment": r.equipment_id,
                "equipment_name": r.equipment.name if r.equipment else None,
                "channel_name": r.channel.name if r.channel else None,
                "work_type": r.work_type,
                "work_type_display": r.get_work_type_display(),
                "priority": r.priority,
                "due_date": r.due_date,
                "rationale": r.rationale,
            }
            for r in recs[:100]
        ],
        "planned": [
            {
                "id": o.pk,
                "number": o.number,
                "title": o.title,
                "status": o.status,
                "work_type": o.work_type,
                "priority": o.priority,
                "due_at": o.due_at,
                "node_name": o.node.name,
                "equipment_name": o.equipment.name if o.equipment else None,
                "assignee_name": o.assignee.get_full_name() if o.assignee else None,
            }
            for o in orders[:300]
        ],
        "weeks": [{"week": w, "orders": n} for w, n in sorted(weeks.items())],
    }


def _due_at(day: date) -> datetime:
    return timezone.make_aware(datetime.combine(day, time(18, 0)))


@transaction.atomic
def schedule_equipment(
    eq: Equipment, day: date, user, work_type: str | None = None, priority: str = "medium"
) -> WorkOrder:
    """Поставить ТО единицы в план: черновик заявки на выбранную дату."""
    work = work_type or work_for(eq)
    due = eq.next_maintenance_at
    basis = (
        f"Плановое ТО по регламенту (раз в {eq.maintenance_interval_days} сут), срок {due:%d.%m.%Y}."
        if due
        else "Профилактические работы."
    )
    if eq.condition:
        basis += f" Фактическое состояние: {eq.get_condition_display().lower()}."
    return WorkOrder.objects.create(
        number=next_number(),
        node=eq.node,
        equipment=eq,
        work_type=work,
        priority=priority,
        title=f"{WorkType(work).label} — {eq.name}"[:255],
        description=f"Основание: план ТО.\n{basis}",
        due_at=_due_at(day),
        created_by=user,
    )


@transaction.atomic
def schedule_recommendation(rec: MaintenanceRecommendation, day: date, user) -> WorkOrder:
    from .services import draft_from_recommendation

    order = draft_from_recommendation(rec, user)
    order.due_at = _due_at(day)
    order.save(update_fields=["due_at", "updated_at"])
    return order


@transaction.atomic
def inspect(
    eq: Equipment,
    user,
    condition: str,
    notes: str = "",
    maintenance: bool = False,
    workorder: WorkOrder | None = None,
    at: datetime | None = None,
) -> EquipmentInspection:
    """Осмотр или ТО: фактическое состояние в реестр, при ТО — новая дата последнего обслуживания."""
    at = at or timezone.now()
    inspection = EquipmentInspection.objects.create(
        equipment=eq,
        inspected_at=at,
        inspector=user,
        condition=condition,
        notes=notes,
        maintenance=maintenance,
        workorder=workorder,
    )
    eq.condition, eq.condition_at = condition, at
    fields = ["condition", "condition_at", "updated_at"]
    if maintenance:
        eq.last_maintenance_at = timezone.localtime(at).date()
        fields.append("last_maintenance_at")
    eq.save(update_fields=fields)
    # Исправная единица снимает открытые рекомендации «по состоянию»
    if condition in (EquipmentCondition.GOOD, EquipmentCondition.REMARKS):
        MaintenanceRecommendation.objects.filter(
            equipment=eq,
            status=MaintenanceRecommendation.Status.NEW,
            rationale__startswith="По последнему осмотру",
        ).update(status=MaintenanceRecommendation.Status.DONE)
    return inspection


@transaction.atomic
def complete(order: WorkOrder, user, report: str = "", condition: str = "", notes: str = "") -> WorkOrder:
    """Бригада отмечает выполнение: отчёт, фактическое состояние оборудования, ТО в реестре."""
    order.status = WorkOrder.Status.DONE
    if report:
        order.report = report
    order.save()
    if order.equipment_id and condition:
        inspect(
            order.equipment,
            user,
            condition,
            notes or report,
            maintenance=order.work_type in MAINTENANCE_WORK,
            workorder=order,
        )
    if order.recommendation_id:
        MaintenanceRecommendation.objects.filter(pk=order.recommendation_id).update(
            status=MaintenanceRecommendation.Status.DONE
        )
    if order.incident_id:
        state = (
            f" Состояние оборудования: {EquipmentCondition(condition).label.lower()}." if condition else ""
        )
        IncidentEvent.objects.create(
            incident_id=order.incident_id,
            kind=IncidentEvent.Kind.WORKORDER,
            actor=user,
            text=f"Заявка {order.number} выполнена: {report or 'без отчёта'}.{state}"[:512],
            payload={"workorder_id": order.pk},
        )
    return order


def export_xlsx(data: dict) -> bytes:
    """План ТО зоны в XLSX: регламентные работы, рекомендации и запланированные заявки."""
    import io

    from openpyxl import Workbook
    from openpyxl.styles import Font

    wb = Workbook()
    sheets = [
        (
            "Регламент",
            [
                "Оборудование",
                "Инв. номер",
                "Объект",
                "Пикет",
                "Последнее ТО",
                "Срок",
                "Просрочено, сут",
                "Состояние",
                "В плане",
            ],
            [
                [
                    r["name"],
                    r["inventory_number"],
                    r["node_name"],
                    r["picket"],
                    r["last_maintenance_at"],
                    r["due"],
                    r["overdue_days"] or None,
                    r["condition_display"],
                    r["planned"]["number"] if r["planned"] else "",
                ]
                for r in data["due"]
            ],
        ),
        (
            "Рекомендации",
            ["Объект", "Оборудование / канал", "Вид работ", "Приоритет", "Срок", "Обоснование"],
            [
                [
                    r["node_name"],
                    r["equipment_name"] or r["channel_name"] or "",
                    r["work_type_display"],
                    r["priority"],
                    r["due_date"],
                    r["rationale"],
                ]
                for r in data["recommendations"]
            ],
        ),
        (
            "Запланировано",
            ["Номер", "Работа", "Объект", "Оборудование", "Срок", "Статус", "Исполнитель"],
            [
                [
                    r["number"],
                    r["title"],
                    r["node_name"],
                    r["equipment_name"] or "",
                    timezone.localtime(r["due_at"]).replace(tzinfo=None),
                    WorkOrder.Status(r["status"]).label,
                    r["assignee_name"] or "",
                ]
                for r in data["planned"]
            ],
        ),
    ]
    wb.remove(wb.active)
    for title, header, rows in sheets:
        ws = wb.create_sheet(title)
        ws.append(header)
        for cell in ws[1]:
            cell.font = Font(bold=True)
        for row in rows:
            ws.append(row)
        for column in ws.columns:
            ws.column_dimensions[column[0].column_letter].width = min(
                60, max(12, *(len(str(c.value or "")) for c in column))
            )
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
