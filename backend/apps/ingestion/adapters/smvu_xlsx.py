from __future__ import annotations

from collections.abc import Iterator
from datetime import date, datetime, time
from pathlib import Path

from openpyxl import load_workbook

from .base import RawEvent
from .smvu_csv import MSK, parse_bool

REQUIRED = ("ид_события", "ид_канала_данных", "дата", "время", "тревожное", "значение_датчика")


def _as_datetime(day, moment) -> datetime:
    """Excel хранит дату и время либо типизированно, либо строкой — поддерживаем оба варианта."""
    if isinstance(day, datetime):
        day = day.date()
    if not isinstance(day, date):
        day = date.fromisoformat(str(day).strip())
    if isinstance(moment, datetime):
        moment = moment.time()
    if not isinstance(moment, time):
        moment = time.fromisoformat(str(moment).strip())
    return datetime.combine(day, moment, tzinfo=MSK)


class SmvuXlsxAdapter:
    """Тот же журнал сработок, выгруженный в XLSX (ТЗ §7: файловый обмен CSV, XLSX)."""

    key = "smvu_xlsx"

    def iter_events(self, path: str | Path, sheet: str | None = None, **_) -> Iterator[RawEvent]:
        workbook = load_workbook(path, read_only=True, data_only=True)
        worksheet = workbook[sheet] if sheet else workbook.active
        rows = worksheet.iter_rows(values_only=True)
        header = [str(h).strip() if h is not None else "" for h in next(rows, ())]
        missing = [c for c in REQUIRED if c not in header]
        if missing:
            raise ValueError(f"В файле нет колонок: {', '.join(missing)}")
        idx = {name: header.index(name) for name in REQUIRED}
        try:
            for row in rows:
                try:
                    yield RawEvent(
                        event_id=int(row[idx["ид_события"]]),
                        channel_external_id=int(row[idx["ид_канала_данных"]]),
                        ts=_as_datetime(row[idx["дата"]], row[idx["время"]]),
                        raw_value=""
                        if row[idx["значение_датчика"]] is None
                        else str(row[idx["значение_датчика"]]),
                        raw_alarm=parse_bool(str(row[idx["тревожное"]] or "")),
                    )
                except (TypeError, ValueError):
                    continue
        finally:
            workbook.close()
