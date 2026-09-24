import logging
import uuid
from pathlib import Path

import polars as pl
from celery import shared_task

from . import archive, history
from .adapters import REGISTRY
from .models import ImportJob
from .pipeline import Pipeline

logger = logging.getLogger(__name__)

BATCH = 5000


@shared_task(acks_late=True)
def run_history_batch(batch: str) -> str:
    """Импорт истории, запущенный из интерфейса. Прогресс — в заданиях пакета."""
    history.run_batch(uuid.UUID(batch))
    return batch


@shared_task(acks_late=True)
def run_import_job(job_id: int) -> dict:
    """
    Импорт загруженного файла журнала. CSV идёт векторным путём архива (миллионы строк),
    XLSX — построчным конвейером потока (Excel-выгрузки заведомо небольшие).
    """
    job = ImportJob.objects.select_related("source").get(pk=job_id)
    path = Path(job.file.path if job.file else job.file_path)
    reporter = history.Reporter(job, None)
    history.start_job(job)
    try:
        if path.suffix.lower() == ".csv":
            _import_csv(job, path, reporter)
        else:
            _import_rows(job, path, reporter)
        history.finish_job(job)
    except Exception as exc:
        logger.exception("import job %s failed", job_id)
        history.finish_job(job, exc)
    return {"job": job.pk, "status": job.status, "rows_ok": job.rows_ok}


def _import_csv(job: ImportJob, path: Path, reporter) -> None:
    parquet = history.archive_dir() / f"upload_{job.pk}.parquet"
    report = archive.normalize_journal(path, parquet, progress=lambda f, s: reporter(f * 0.6, s))
    reporter(0.62, "Суточная витрина")
    report["daily_rows"] = archive.load_daily(archive.daily_frame(pl.scan_parquet(parquet)))
    # Загруженный файл — новые данные: в оперативный контур идёт целиком
    since = pl.scan_parquet(parquet).select(pl.col("ts").min()).collect().item()
    if since is not None:
        archive.load_readings([parquet], since, progress=lambda f, s: reporter(0.65 + f * 0.35, s))
    job.quality = report
    job.rows_total, job.rows_ok = report["rows_total"], report["rows_loaded"]
    job.rows_skipped = report["rows_total"] - report["rows_loaded"]


def _import_rows(job: ImportJob, path: Path, reporter) -> None:
    from openpyxl import load_workbook

    workbook = load_workbook(path, read_only=True)
    total = max((workbook.active.max_row or 1) - 1, 1)
    workbook.close()
    pipeline = Pipeline()
    adapter = REGISTRY["smvu_xlsx" if path.suffix.lower() in (".xlsx", ".xlsm") else job.source.adapter]
    batch = []
    for event in adapter.iter_events(path=path):
        batch.append(event)
        if len(batch) >= BATCH:
            _flush(job, pipeline, batch)
            reporter(job.rows_total / total, f"Обработано {job.rows_total:,} из ~{total:,} строк")
            batch = []
    _flush(job, pipeline, batch)


def _flush(job: ImportJob, pipeline: Pipeline, batch: list) -> None:
    if not batch:
        return
    result = pipeline.process(batch)
    job.rows_total += result.received
    job.rows_ok += result.stored
    job.rows_skipped += result.skipped_unknown_channel
    quality = job.quality or {"by_quality": {}, "by_state": {}}
    for key, counter in (("by_quality", result.by_quality), ("by_state", result.by_state)):
        for name, count in counter.items():
            quality[key][name] = quality[key].get(name, 0) + count
    job.quality = quality
    job.save(update_fields=["rows_total", "rows_ok", "rows_skipped", "quality"])
