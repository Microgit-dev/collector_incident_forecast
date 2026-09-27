"""
Синхронизация с реестром оборудования заказчика (ТЗ §6, §10, §13): по API учётной системы
(REGISTRY_MODE=api, раз в сутки) или выгрузкой CSV/XLSX (REGISTRY_MODE=file, раздел «Интеграции»).

Формат строки (одинаков для API и файла, лишние поля игнорируются):
    inventory_number  инвентарный номер — ключ записи (обязателен)
    name              наименование (обязателен)
    kind              вид: hatch, door, vent_shaft, chamber, pump, fan, ups, cabinet, sensor
                      или русское название («Насос», «Люк»…)
    object_id         ид_объект из справочника объектов (или object — диспетчерское название)
    picket            пикет
    commissioned_at   дата ввода в эксплуатацию (ГГГГ-ММ-ДД или ДД.ММ.ГГГГ)
    last_maintenance_at  дата последнего ТО
    maintenance_interval_days  регламентный интервал ТО, сут
    mtbf_hours        средняя наработка на отказ, ч
    channel_ids       ид_канала_данных через запятую — каналы СМВУ этой единицы
    is_active         0/нет — списано

Записи реестра получают source=imported; ручные записи (manual) не трогаются; эмулированные
записи (emulated) при первой успешной синхронизации удаляются — их заменяет реестр заказчика.
Единицы, пропавшие из полного реестра, помечаются выведенными из эксплуатации, а не удаляются:
на них ссылаются заявки и осмотры.
"""

from __future__ import annotations

import csv
import io
from datetime import date, datetime

from django.db import transaction
from django.utils import timezone

from apps.topology.models import Node

from .models import Channel, Equipment, EquipmentKind

KIND_BY_LABEL = {label.lower(): value for value, label in EquipmentKind.choices}
REQUIRED = ("inventory_number", "name")


class RegistryError(ValueError):
    pass


def _date(value) -> date | None:
    if value in (None, ""):
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = str(value).strip()
    for fmt in ("%Y-%m-%d", "%d.%m.%Y"):
        try:
            return datetime.strptime(text[:10], fmt).date()
        except ValueError:
            continue
    raise RegistryError(f"неверная дата «{text}»")


def _int(value) -> int | None:
    if value in (None, ""):
        return None
    return int(float(str(value).replace(",", ".")))


def _bool(value) -> bool:
    return str(value).strip().lower() not in {"0", "false", "нет", "no", "списано"}


def _kind(value) -> str:
    text = str(value or "").strip().lower()
    if text in EquipmentKind.values:
        return text
    if text in KIND_BY_LABEL:
        return KIND_BY_LABEL[text]
    raise RegistryError(f"неизвестный вид оборудования «{value}»")


def read_file(name: str, content: bytes) -> list[dict]:
    """Строки выгрузки CSV (UTF-8 или cp1251, разделитель , или ;) или XLSX (первый лист)."""
    if name.lower().endswith(".xlsx"):
        from openpyxl import load_workbook

        sheet = load_workbook(io.BytesIO(content), read_only=True, data_only=True).active
        rows = sheet.iter_rows(values_only=True)
        header = [str(h or "").strip() for h in next(rows, [])]
        return [dict(zip(header, r, strict=False)) for r in rows if any(v not in (None, "") for v in r)]
    try:
        text = content.decode("utf-8-sig")
    except UnicodeDecodeError:
        text = content.decode("cp1251")
    dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;")
    return [dict(r) for r in csv.DictReader(io.StringIO(text), dialect=dialect)]


@transaction.atomic
def apply(rows: list[dict], full: bool = True) -> dict:
    """
    Записать строки реестра. full=True — это полный реестр: единицы, которых в нём нет,
    выводятся из эксплуатации. Ошибки строк не останавливают загрузку, а возвращаются списком.
    """
    nodes_by_ext = dict(Node.objects.exclude(external_id=None).values_list("external_id", "pk"))
    nodes_by_name = {}
    for pk, name in Node.objects.order_by("-depth").values_list("pk", "name"):
        nodes_by_name[name.strip().lower()] = pk  # при совпадении имён — самый верхний узел
    channels = dict(Channel.objects.values_list("external_id", "pk"))
    now = timezone.now()
    seen: set[str] = set()
    created = updated = 0
    errors: list[dict] = []
    for i, raw in enumerate(rows, start=1):
        row = {str(k).strip(): v for k, v in raw.items() if k}
        try:
            missing = [f for f in REQUIRED if not row.get(f)]
            if missing:
                raise RegistryError(f"нет обязательных полей: {', '.join(missing)}")
            number = str(row["inventory_number"]).strip()
            if (ext := _int(row.get("object_id"))) is not None:
                node_id = nodes_by_ext.get(ext)
            else:
                node_id = nodes_by_name.get(str(row.get("object") or "").strip().lower())
            if node_id is None:
                raise RegistryError(f"объект не найден: {row.get('object_id') or row.get('object') or '—'}")
            fields = {
                "kind": _kind(row.get("kind")),
                "node_id": node_id,
                "name": str(row["name"]).strip()[:255],
                "picket": float(str(row["picket"]).replace(",", "."))
                if row.get("picket") not in (None, "")
                else None,
                "commissioned_at": _date(row.get("commissioned_at")),
                "last_maintenance_at": _date(row.get("last_maintenance_at")),
                "maintenance_interval_days": _int(row.get("maintenance_interval_days")),
                "mtbf_hours": _int(row.get("mtbf_hours")),
                "is_active": _bool(row.get("is_active", "1")),
                "source": "imported",
                "synced_at": now,
            }
        except (RegistryError, ValueError, TypeError) as exc:
            errors.append({"row": i, "inventory_number": row.get("inventory_number"), "error": str(exc)})
            continue
        existing = Equipment.objects.filter(inventory_number=number).exclude(source="manual").first()
        if existing is None and Equipment.objects.filter(inventory_number=number, source="manual").exists():
            errors.append(
                {"row": i, "inventory_number": number, "error": "запись ведётся вручную, не изменена"}
            )
            seen.add(number)
            continue
        if existing:
            # дата ТО из реестра не откатывает более позднее ТО, отмеченное бригадой в системе
            if existing.last_maintenance_at and fields["last_maintenance_at"]:
                fields["last_maintenance_at"] = max(
                    existing.last_maintenance_at, fields["last_maintenance_at"]
                )
            for key, value in fields.items():
                setattr(existing, key, value)
            existing.save()
            eq, is_new = existing, False
        else:
            eq, is_new = Equipment.objects.create(inventory_number=number, **fields), True
        if ids := row.get("channel_ids"):
            wanted = [
                channels[c]
                for c in (_int(x) for x in str(ids).replace(";", ",").split(",") if x.strip())
                if c in channels
            ]
            eq.channels.set(wanted)
        created += is_new
        updated += not is_new
        seen.add(number)
    retired = removed = 0
    if seen:
        removed, _ = Equipment.objects.filter(source="emulated").delete()
        if full:
            retired = (
                Equipment.objects.filter(source="imported", is_active=True)
                .exclude(inventory_number__in=seen)
                .update(is_active=False, synced_at=now)
            )
    return {
        "rows": len(rows),
        "created": created,
        "updated": updated,
        "retired": retired,
        "emulated_removed": removed,
        "errors": errors[:200],
        "error_count": len(errors),
    }


def sync_from_api() -> dict:
    from apps.integrations.clients import RegistryClient

    return apply(RegistryClient().rows(), full=True)


def template_csv() -> str:
    """Шаблон выгрузки для заказчика: заголовок и пример строки."""
    header = [
        "inventory_number",
        "name",
        "kind",
        "object_id",
        "object",
        "picket",
        "commissioned_at",
        "last_maintenance_at",
        "maintenance_interval_days",
        "mtbf_hours",
        "channel_ids",
        "is_active",
    ]
    example = [
        "НС-0001",
        "Насос дренажный АНС-1",
        "Насос",
        "",
        "объект Мю",
        "12.5",
        "2018-05-01",
        "2026-03-10",
        "90",
        "20000",
        "",
        "1",
    ]
    return ";".join(header) + "\n" + ";".join(example) + "\n"
