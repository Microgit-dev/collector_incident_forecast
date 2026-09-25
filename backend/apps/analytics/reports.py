"""
Отчёты для руководства (ТЗ §8): PDF — сводка за период, XLSX — та же сводка и журналы
инцидентов и прогнозов для самостоятельного анализа. Данные — те же функции, что на странице
«Аналитика», поэтому цифры в отчёте и на экране совпадают.
"""

from __future__ import annotations

import io
from collections import Counter
from datetime import datetime
from pathlib import Path
from zoneinfo import ZoneInfo

from django.utils import timezone

from apps.forecasting.models import Prediction
from apps.incidents.models import IncidentType
from apps.topology.selectors import scope_queryset, user_scope_node

from .efficiency import efficiency, incidents_in
from .quality import quality
from .staff import staff_metrics

MSK = ZoneInfo("Europe/Moscow")
TASK_TITLE = {
    "sensor_failure": "Отказ датчика",
    "gas": "Загазованность",
    "flood": "Подтопление",
    "fire": "Пожар (индикатор)",
    "intrusion": "НСД (индикатор)",
}
SEVERITY = {"critical": "Критический", "high": "Высокий", "medium": "Средний", "low": "Низкий"}
FONTS = [
    (
        "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
        "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf",
    ),
    ("C:/Windows/Fonts/arial.ttf", "C:/Windows/Fonts/arialbd.ttf"),
]


def collect(user, since: datetime, until: datetime, include_emulated: bool, backtest: bool) -> dict:
    incidents = incidents_in(user, since, until, include_emulated)
    by_type = Counter(incidents.values_list("type", flat=True))
    scope = user_scope_node(user)
    return {
        "since": since,
        "until": until,
        "scope": scope.name if scope else "все объекты",
        "include_emulated": include_emulated,
        "backtest": backtest,
        "by_type": {IncidentType(t).label: n for t, n in by_type.most_common()},
        "efficiency": efficiency(user, since, until, include_emulated),
        "quality": quality(user, since, until, backtest),
        "staff": staff_metrics(user, since, until, include_emulated),
    }


def _pct(v) -> str:
    return "—" if v is None else f"{v * 100:.0f} %"


def _min(v) -> str:
    return "—" if v is None else f"{v:.1f} мин"


def _d(ts: datetime) -> str:
    return ts.astimezone(MSK).strftime("%d.%m.%Y")


def summary_rows(data: dict) -> list[tuple[str, str]]:
    s = data["efficiency"]["summary"]
    return [
        ("Карточек за период", str(s["cards"])),
        ("Карточек на 12-часовую смену", str(s["cards_per_shift"])),
        ("Сигналов потока в карточках-фактах", str(s["signals"])),
        ("Сжатие потока (сигналов на карточку)", f"×{s['reduction']}" if s["reduction"] else "—"),
        ("До просмотра, медиана / 90 %", f"{_min(s['view']['median'])} / {_min(s['view']['p90'])}"),
        ("До взятия в работу, медиана", _min(s["take"]["median"])),
        ("До решения, медиана / 90 %", f"{_min(s['decision']['median'])} / {_min(s['decision']['p90'])}"),
        ("Эскалаций по таймауту", _pct(s["escalated_share"])),
        ("Карточек с решением", _pct(s["decided_share"])),
        ("Закрытых с указанным результатом", _pct(s["closed_with_result_share"])),
    ]


def staff_rows(data: dict) -> list[list]:
    return [
        [
            p["name"],
            p["team"] or "",
            p["decisions"],
            p["view"],
            p["decision"],
            p["escalated_share"],
            p["with_cause_share"],
            p["false_alarm_share"],
        ]
        for p in data["efficiency"]["staff"]
    ]


def first_rows(data: dict) -> list[list]:
    """Кто первый: отклики, гонки, скорость, качество решений, нагрузка на смену, обучение."""
    return [
        [
            p["rank"] or "",
            p["name"],
            p["team"] or "",
            p["zone_cards"],
            p["first_seen"],
            p["responded"],
            p["responded_share"],
            p["response_median"],
            f"{p['races_won']} из {p['races']}" if p["races"] else "—",
            p["decision_median"],
            p["quality"],
            p["takeovers_lost"],
            p["per_shift"],
            p["training_done"],
        ]
        for p in data["staff"]["people"]
        if p["active"]
    ]


def team_rows(data: dict) -> list[list]:
    return [
        [
            t["team"],
            t["members"],
            t["zone_cards"],
            t["responded"],
            t["responded_share"],
            t["response_median"],
            t["escalated_unanswered"],
            t["quality"],
            t["training_done"],
        ]
        for t in data["staff"]["teams"]
    ]


def quality_rows(data: dict) -> list[list]:
    out = []
    for task, q in data["quality"]["tasks"].items():
        recall = q["recall"] or {}
        model = q["model"] or {}
        out.append(
            [
                TASK_TITLE[task],
                q["total"],
                q["per_day"],
                q["precision"],
                q["false_share"],
                recall.get("warned"),
                recall.get("onsets"),
                recall.get("recall"),
                model.get("test_precision_high"),
                model.get("lead_time_median_h"),
            ]
        )
    return out


# ---------------- PDF ----------------


def to_pdf(data: dict) -> bytes:
    from fpdf import FPDF
    from fpdf.fonts import FontFace

    regular, bold = next(((r, b) for r, b in FONTS if Path(r).exists()), (None, None))
    if regular is None:
        raise RuntimeError("Нет шрифта с кириллицей для PDF (установите fonts-dejavu-core)")
    pdf = FPDF(orientation="P", unit="mm", format="A4")
    pdf.add_font("main", "", regular)
    pdf.add_font("main", "B", bold if Path(bold).exists() else regular)
    pdf.set_auto_page_break(True, margin=15)
    pdf.add_page()
    heading = FontFace(emphasis="BOLD", fill_color=(235, 238, 245))

    def h1(text):
        pdf.set_font("main", "B", 16)
        pdf.multi_cell(0, 8, text, new_x="LMARGIN", new_y="NEXT")

    def h2(text):
        pdf.ln(3)
        pdf.set_font("main", "B", 12)
        pdf.multi_cell(0, 7, text, new_x="LMARGIN", new_y="NEXT")
        pdf.set_font("main", "", 9)

    def para(text, size=9):
        pdf.set_font("main", "", size)
        pdf.multi_cell(0, 5, text, new_x="LMARGIN", new_y="NEXT")

    def table(header, rows, widths):
        pdf.set_font("main", "", 8.5)
        with pdf.table(col_widths=widths, headings_style=heading, line_height=5, text_align="LEFT") as t:
            r = t.row()
            for h in header:
                r.cell(h)
            for row in rows:
                r = t.row()
                for v in row:
                    r.cell("—" if v is None or v == "" else str(v))

    h1("Отчёт о работе системы прогнозирования инцидентов")
    para(
        f"Период: {_d(data['since'])} — {_d(data['until'])} · зона: {data['scope']} · "
        f"сформирован {timezone.localtime().strftime('%d.%m.%Y %H:%M')}"
    )
    emulated = data["efficiency"]["emulated"]
    if emulated:
        pdf.set_text_color(170, 60, 0)
        para(
            f"Внимание: {emulated} карточек из эмуляции смен (демонстрация на исторических данных). "
            "Сформируйте отчёт без эмуляции, чтобы видеть только реальную работу."
        )
        pdf.set_text_color(0, 0, 0)

    h2("1. Нагрузка и работа диспетчеров")
    table(["Показатель", "Значение"], summary_rows(data), (110, 70))
    if data["by_type"]:
        h2("Карточки по типу угрозы")
        table(["Тип", "Карточек"], list(data["by_type"].items()), (110, 70))
    sev = data["efficiency"]["by_severity"]
    if sev:
        h2("Скорость реакции по уровню")
        table(
            ["Уровень", "Карточек", "До просмотра, медиана", "До решения, медиана", "Эскалаций"],
            [
                [
                    SEVERITY[k],
                    v["cards"],
                    _min(v["view"]["median"]),
                    _min(v["decision"]["median"]),
                    _pct(v["escalated_share"]),
                ]
                for k, v in sev.items()
            ],
            (30, 25, 45, 45, 30),
        )
    first = first_rows(data)
    if first:
        h2("2. Сотрудники: кто первый откликнулся")
        st = data["staff"]["summary"]
        para(
            f"Правило смены: карточку забирает тот, кто первым откликнулся. Гонок (карточку открыли двое и больше "
            f"до отклика) — {st['contested']} ({_pct(st['contested_share'])}), перехватов руководителем — "
            f"{st['takeovers']}, эскалаций без отклика — {st['escalated_unanswered']}. Качество — доля закрытий "
            "как ложных (пожар, газ, вода, проникновение), после которых угроза не повторилась на объекте за 6 часов.",
            8,
        )
        table(
            ["№", "Сотрудник", "Первым", "Доля зоны", "Отклик", "Гонки", "Решение", "Качество", "На смену"],
            [
                [r or "", n, rs, _pct(sh), _min(rm), races, _min(dm), _pct(q), ps if ps is not None else "—"]
                for r, n, _t, _z, _fs, rs, sh, rm, races, dm, q, _tl, ps, _tr in first
            ],
            (10, 40, 16, 20, 22, 20, 22, 20, 18),
        )
        teams = team_rows(data)
        if teams:
            table(
                [
                    "Команда",
                    "Чел.",
                    "Карточек зоны",
                    "Первым",
                    "Доля",
                    "Отклик",
                    "Эскал. без отклика",
                    "Качество",
                ],
                [[t, m, z, r, _pct(sh), _min(rm), e, _pct(q)] for t, m, z, r, sh, rm, e, q, _ in teams],
                (46, 12, 24, 18, 16, 22, 26, 20),
            )
    if data["efficiency"]["staff"]:
        h2("2б. Решения сотрудников")
        table(
            ["Сотрудник", "Команда", "Решений", "Просмотр", "Решение", "Эскал.", "С причиной", "Ложных"],
            [
                [n, t, d, _min(v), _min(dm), _pct(e), _pct(c), _pct(f)]
                for n, t, d, v, dm, e, c, f in staff_rows(data)
            ],
            (32, 34, 20, 20, 20, 16, 20, 16),
        )
    h2("3. Качество прогнозов" + (" (бэктест)" if data["backtest"] else ""))
    table(
        [
            "Задача",
            "Записей",
            "В сутки",
            "Точность",
            "Ложных",
            "Предупр.",
            "Полнота",
            "Точн. тест",
            "Упрежд.",
        ],
        [
            [
                t,
                n,
                pd,
                _pct(p),
                _pct(f),
                f"{w} из {o}" if o else "—",
                _pct(r),
                _pct(tp),
                f"{lt} ч" if lt else "—",
            ]
            for t, n, pd, p, f, w, o, r, tp, lt in quality_rows(data)
        ],
        (30, 17, 16, 18, 16, 27, 18, 20, 18),
    )
    para(
        "Точность — доля подтвердившихся среди прогнозов уровня «высокий» и выше с известным исходом. "
        "Полнота — доля начавшихся за период событий, о которых заранее предупредил прогноз такого уровня. "
        "Для пожара и НСД — индекс по правилам, исход — «угроза проявилась» (повторные тревоги).",
        8,
    )
    fc = data["quality"]["forecast_cards"]
    if fc["decided"]:
        u = fc["useful"]
        para(
            f"Прогнозные карточки с решением: {fc['decided']}; прогноз помог — {u.get('yes', 0)}, "
            f"не помог — {u.get('no', 0)}, не ясно — {u.get('unknown', 0)}."
        )
    causes = data["efficiency"]["causes"]
    if causes:
        h2("4. Что происходило на самом деле (по решениям диспетчеров)")
        table(["Что произошло", "Решений"], [[c["title"], c["count"]] for c in causes], (110, 70))
    return bytes(pdf.output())


# ---------------- XLSX ----------------


def to_xlsx(data: dict, user) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    bold = Font(bold=True)
    fill = PatternFill("solid", fgColor="EBEEF5")

    def sheet(title, header, rows, first=False):
        ws = wb.active if first else wb.create_sheet()
        ws.title = title
        ws.append(header)
        for cell in ws[1]:
            cell.font, cell.fill = bold, fill
        for row in rows:
            ws.append(list(row))
        for i, h in enumerate(header, 1):
            width = max([len(str(h))] + [len(str(r[i - 1])) for r in rows[:500] if r[i - 1] is not None])
            ws.column_dimensions[get_column_letter(i)].width = min(max(width + 2, 10), 60)
        ws.freeze_panes = "A2"
        return ws

    meta = [
        ("Период", f"{_d(data['since'])} — {_d(data['until'])}"),
        ("Зона", data["scope"]),
        ("Эмулированных карточек", data["efficiency"]["emulated"]),
        ("Журнал прогнозов", "бэктест" if data["backtest"] else "оперативный"),
        *summary_rows(data),
    ]
    sheet("Сводка", ["Показатель", "Значение"], meta, first=True)
    sheet(
        "Кто первый",
        [
            "Место",
            "Сотрудник",
            "Команда",
            "Карточек зоны",
            "Первым заметил",
            "Первым откликнулся",
            "Доля карточек зоны",
            "До отклика, медиана мин",
            "Гонки: выиграно",
            "От отклика до решения, мин",
            "Качество решений",
            "Перехвачено руководителем",
            "Откликов на смену",
            "Учебных заданий",
        ],
        first_rows(data),
    )
    sheet(
        "Команды",
        [
            "Команда",
            "Сотрудников",
            "Карточек зоны",
            "Откликнулись первыми",
            "Доля",
            "До отклика, медиана мин",
            "Эскалаций без отклика",
            "Качество решений",
            "Учебных заданий",
        ],
        team_rows(data),
    )
    sheet(
        "Решения сотрудников",
        [
            "Сотрудник",
            "Команда",
            "Решений",
            "До просмотра, мин",
            "До решения, мин",
            "Доля эскалаций",
            "Доля с причиной",
            "Доля ложных",
        ],
        staff_rows(data),
    )
    sheet(
        "Качество прогнозов",
        [
            "Задача",
            "Записей",
            "В сутки",
            "Точность",
            "Доля ложных",
            "Предупреждено событий",
            "Событий",
            "Полнота",
            "Точность на тесте (высокий)",
            "Упреждение, ч",
        ],
        quality_rows(data),
    )
    sheet(
        "По суткам",
        ["Сутки", "Карточек", "Эскалаций", "До решения, медиана мин"],
        [[d["day"], d["cards"], d["escalated"], d["decision_median"]] for d in data["efficiency"]["daily"]],
    )

    def local(ts):
        return ts.astimezone(MSK).replace(tzinfo=None) if ts else None

    incidents = incidents_in(user, data["since"], data["until"], data["include_emulated"]).select_related(
        "node", "assigned_to"
    )
    sheet(
        "Журнал инцидентов",
        [
            "№",
            "Открыт",
            "Тип",
            "Уровень",
            "Статус",
            "Объект",
            "Прогноз",
            "Сигналов",
            "Каналов",
            "Приоритет",
            "Эскалация",
            "Ответственный",
            "Решён",
            "Эмуляция",
        ],
        [
            [
                i.pk,
                local(i.opened_at),
                i.get_type_display(),
                i.get_severity_display(),
                i.get_status_display(),
                i.node.name,
                "да" if i.is_forecast else "",
                i.signals_count,
                i.channels_count,
                i.priority,
                i.escalation_level,
                i.assigned_to.get_full_name() if i.assigned_to else "",
                local(i.resolved_at),
                "да" if i.is_emulated else "",
            ]
            for i in incidents.order_by("opened_at")[:20000]
        ],
    )
    predictions = scope_queryset(
        Prediction.objects.filter(
            issued_at__gte=data["since"], issued_at__lt=data["until"], is_backtest=data["backtest"]
        ),
        user,
        "node",
    ).select_related("node", "channel")
    sheet(
        "Журнал прогнозов",
        [
            "№",
            "Сформирован",
            "Задача",
            "Объект",
            "Канал",
            "Вероятность / индекс",
            "Уровень",
            "Действует до",
            "Исход",
        ],
        [
            [
                p.pk,
                local(p.issued_at),
                TASK_TITLE.get(p.task, p.task),
                p.node.name,
                p.channel.name if p.channel else "",
                round(p.probability, 4),
                p.get_risk_level_display(),
                local(p.valid_until),
                p.get_outcome_display(),
            ]
            for p in predictions.order_by("issued_at")[:50000]
        ],
    )
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
