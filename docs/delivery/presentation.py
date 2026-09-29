"""
Презентация решения на шаблоне конкурса «Лидеры цифровой трансформации 2026» (data/ЛЦТ2026 Шаблон
презентации.pptx) → docs/delivery/out/Презентация.pptx.

Обязательные слайды шаблона 7–11 заполняются в исходных фигурах (дизайн и сетка не меняются, на слайде
команды лишние карточки удаляются). Описательная часть — свои слайды на макетах шаблона: белые
карточки, розовые плашки заголовков, BPMN-схемы процессов и схема ML-конвейера нативными фигурами
PowerPoint (их можно править).

Данные команды — в TEAM ниже (заполняются перед сдачей).

    uv run --no-project --with python-pptx --with pywin32 python docs/delivery/presentation.py
"""

from __future__ import annotations

import copy
from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Cm, Pt

HERE = Path(__file__).resolve().parent
ROOT = HERE.parents[1]
TEMPLATE = ROOT / "data" / "ЛЦТ2026 Шаблон презентации.pptx"
OUT = HERE / "out" / "Презентация.pptx"
IMG = HERE / "img"
LOGO_DZHKH = "Image 8"  # «Департамент жилищно-коммунального хозяйства города Москвы» со слайда 6 шаблона

PINK = RGBColor(0xFF, 0x00, 0x53)
PINK_LIGHT = RGBColor(0xFF, 0xD6, 0xE4)
LAVENDER = RGBColor(0x8A, 0x83, 0xD1)
DARK = RGBColor(0x31, 0x0F, 0x53)
PURPLE = RGBColor(0x52, 0x09, 0x78)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
INK = RGBColor(0x1C, 0x1D, 0x22)
GREY = RGBColor(0x6B, 0x6B, 0x7B)
LANE = RGBColor(0xF3, 0xF0, 0xFA)
FONT = "Montserrat"

# ---------------------------------------------------------------- данные команды (заполнить перед сдачей)
TEAM = {
    "name": "[Название команды]",
    "captain": "[ФИО капитана, специальность]",
    "size": "[N]",
    "about": "[как образовалась команда; место работы или учёбы участников]",
    "city": "[город и регион]",
    "history": "[Краткая история команды: как собрались, участвовали ли вместе в хакатонах или проектах]",
    "members": [
        # {"name": "Имя Фамилия", "role": "...", "nick": "@...", "phone": "+7 ...", "work": "..."},
        {"name": "[Имя Фамилия]", "role": "[роль в команде]", "nick": "[ник]", "phone": "[телефон]", "work": "[место работы/учёбы]"},
    ],
}

# ---------------------------------------------------------------- базовые элементы


def _font(run, size=None, bold=None, color=None, name=FONT):
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.font.bold = bold
    if color is not None:
        run.font.color.rgb = color
    run.font.name = name


def text(slide, x, y, w, h, lines, size=12, color=INK, bold=False, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, spacing=1.0):
    """Текстовое поле. lines — строки; «• » в начале — пункт списка; («текст», {"bold": True, …}) — стиль строки."""
    box = slide.shapes.add_textbox(Cm(x), Cm(y), Cm(w), Cm(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Cm(0.1)
    tf.margin_top = tf.margin_bottom = Cm(0.05)
    for i, line in enumerate(lines):
        style = {}
        if isinstance(line, tuple):
            line, style = line
        para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        para.alignment = style.get("align", align)
        para.line_spacing = spacing
        if style.get("space"):
            para.space_before = Pt(style["space"])
        run = para.add_run()
        run.text = line
        _font(run, style.get("size", size), style.get("bold", bold), style.get("color", color))
    return box


def rect(slide, x, y, w, h, fill=WHITE, line=None, radius=0.08, shape=MSO_SHAPE.ROUNDED_RECTANGLE, line_width=1.0):
    s = slide.shapes.add_shape(shape, Cm(x), Cm(y), Cm(w), Cm(h))
    if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        s.adjustments[0] = radius
    if fill is None:
        s.fill.background()
    else:
        s.fill.solid()
        s.fill.fore_color.rgb = fill
    if line is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = line
        s.line.width = Pt(line_width)
    s.shadow.inherit = False
    s.text_frame.margin_left = s.text_frame.margin_right = Cm(0.15)
    return s


def label(shape, lines, size=11, color=INK, bold=False, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE):
    tf = shape.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    for i, line in enumerate(lines):
        style = {}
        if isinstance(line, tuple):
            line, style = line
        para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        para.alignment = style.get("align", align)
        run = para.add_run()
        run.text = line
        _font(run, style.get("size", size), style.get("bold", bold), style.get("color", color))
    return shape


def header(slide, title, color=PINK):
    """Плашка заголовка слайда, как в шаблоне: скруглённый прямоугольник, белый заглавный текст."""
    width = min(16.0, 1.6 + 0.42 * len(title))
    pill = rect(slide, 1.0, 0.9, width, 1.5, fill=color, radius=0.3)
    label(pill, [title.upper()], size=16, color=WHITE, bold=True)
    return pill


def card(slide, x, y, w, h, title=None, lines=(), size=12.5, title_color=PURPLE, fill=WHITE):
    rect(slide, x, y, w, h, fill=fill, radius=0.06)
    top = y + 0.35
    if title:
        text(slide, x + 0.4, top, w - 0.8, 1.0, [title], size=14, bold=True, color=title_color)
        top += 1.2
    if lines:
        text(slide, x + 0.4, top, w - 0.8, h - (top - y) - 0.2, list(lines), size=size, color=INK)


def picture(slide, path, x, y, w=None, h=None, frame=True):
    pic = slide.shapes.add_picture(str(path), Cm(x), Cm(y), Cm(w) if w else None, Cm(h) if h else None)
    if frame:
        border = rect(slide, x - 0.15, y - 0.15, pic.width / 360000 + 0.3, pic.height / 360000 + 0.3, fill=WHITE, radius=0.02)
        pic._element.addprevious(border._element)
    return pic


def arrow(slide, a, b, a_side="right", b_side="left", color=PURPLE, width=1.5, elbow=True, dash=False):
    """Стрелка между фигурами (точки соединения: top 0, left 1, bottom 2, right 3)."""
    sides = {"top": 0, "left": 1, "bottom": 2, "right": 3}
    conn = slide.shapes.add_connector(MSO_CONNECTOR.ELBOW if elbow else MSO_CONNECTOR.STRAIGHT, 0, 0, 0, 0)
    conn.begin_connect(a, sides[a_side])
    conn.end_connect(b, sides[b_side])
    conn.line.color.rgb = color
    conn.line.width = Pt(width)
    ln = conn.line._get_or_add_ln()
    if dash:
        prst = etree.SubElement(ln, qn("a:prstDash"))
        prst.set("val", "dash")
    tail = etree.SubElement(ln, qn("a:tailEnd"))
    tail.set("type", "triangle")
    tail.set("w", "med")
    tail.set("len", "med")
    return conn


def new_slide(prs, layout="Пустой"):
    lay = next(lay for lay in prs.slide_layouts if lay.name == layout)
    slide = prs.slides.add_slide(lay)
    for ph in list(slide.placeholders):
        if ph.placeholder_format.type not in (13,):  # оставляем только номер слайда
            ph._element.getparent().remove(ph._element)
    return slide


def delete_slide(prs, slide):
    sld_ids = prs.slides._sldIdLst
    for sld_id in list(sld_ids):
        if prs.part.related_part(sld_id.rId) is slide.part:
            prs.part.drop_rel(sld_id.rId)
            sld_ids.remove(sld_id)
            return


def set_lines(shape, lines):
    """Заменить текст фигуры шаблона, сохранив оформление первого абзаца и первого прогона."""
    tf = shape.text_frame
    first = tf.paragraphs[0]
    template_p = copy.deepcopy(first._p)
    for p in list(tf.paragraphs)[1:]:
        p._p.getparent().remove(p._p)
    for i, line in enumerate(lines):
        bold = None
        if isinstance(line, tuple):
            line, bold = line
        if i == 0:
            p = first
        else:
            new_p = copy.deepcopy(template_p)
            first._p.getparent().append(new_p)
            p = tf.paragraphs[-1]
        runs = p.runs
        if runs:
            runs[0].text = line
            for r in runs[1:]:
                r._r.getparent().remove(r._r)
            if bold is not None:
                runs[0].font.bold = bold
        else:
            run = p.add_run()
            run.text = line


def by_text(slide, startswith):
    for sh in slide.shapes:
        if sh.has_text_frame and sh.text_frame.text.strip().startswith(startswith):
            return sh
    raise KeyError(startswith)


# ---------------------------------------------------------------- обязательные слайды 7–11


def fill_mandatory(prs):
    slides = list(prs.slides)
    title, about, team, story, brief = slides[6:11]

    # 7 — титул: название команды, задача, логотип постановщика
    logo = next(sh for sh in slides[5].shapes if sh.name == LOGO_DZHKH)
    for ph in title.placeholders:
        idx = ph.placeholder_format.idx
        if idx == 0:
            ph.text_frame.text = TEAM["name"]
        elif idx == 12:
            ph.text_frame.text = (
                "Задача 8 ДЖКХ: сервис прогнозирования инцидентов и управления ремонтными работами "
                "инженерных коллекторов Москвы"
            )
        elif idx == 11:
            x, y, w, h = ph.left, ph.top, ph.width, ph.height
            ph._element.getparent().remove(ph._element)
            pic = title.shapes.add_picture(io_bytes(logo.image.blob), x, y + Cm(0.3), height=h - Cm(0.6))
            if pic.width > w:
                ratio = w / pic.width
                pic.width, pic.height = int(pic.width * ratio), int(pic.height * ratio)

    # 8 — о команде, суть и уникальность
    set_lines(by_text(about, "Капитан"), [
        (f"Капитан: {TEAM['captain']}", True),
        (f"Кол-во участников: {TEAM['size']} человек", True),
        ("Краткое описание: ", True),
        (TEAM["about"], False),
        (f"Город и регион: {TEAM['city']}", True),
    ])
    set_lines(by_text(about, "В чем суть"), [
        "Превращает поток тревог СМВУ в понятные риски: прогноз на сутки, причина, действие — "
        "и ведёт решение диспетчера до заявки бригаде и ТО."
    ])
    set_lines(by_text(about, "Что делает ваше"), [
        "• 47 тревог → 1 карточка с причиной",
        "• Прогноз аварий на сутки вперёд",
        "• No-code: датчик без программиста",
        "• Здание и этажи по снимку (ИИ)",
        "• Сигнал → решение → заявка → ТО",
    ])
    for ph in about.placeholders:
        if ph.placeholder_format.idx == 0:
            ph.text_frame.text = "Прогноз инцидентов коллекторов"
            for run in ph.text_frame.paragraphs[0].runs:
                run.font.color.rgb = PURPLE
        elif ph.placeholder_format.idx == 10:
            x, y, w, h = ph.left, ph.top, ph.width, ph.height
            pic = ph.insert_picture(str(IMG / "workspace.png"))
            pic.left, pic.top, pic.width, pic.height = x, y, w, h

    # 9 — карточки участников: лишние удаляются
    members = TEAM["members"]
    names = [sh for sh in team.shapes if sh.has_text_frame and sh.text_frame.text.strip() == "Имя Фамилия"]
    infos = [sh for sh in team.shapes if sh.has_text_frame and sh.text_frame.text.strip().startswith("Роль в команде")]
    cards = [sh for sh in team.shapes if sh.name.startswith("Скругленный прямоугольник") and sh.height > Cm(10)]
    pics = [sh for sh in team.placeholders if sh.placeholder_format.type == 18]
    key = lambda s: s.left  # noqa: E731
    names, infos, cards, pics = (sorted(v, key=key) for v in (names, infos, cards, pics))
    for i in range(len(cards)):
        if i < len(members):
            m = members[i]
            set_lines(names[i], [m["name"]])
            set_lines(infos[i], [m["role"], m["nick"], m["phone"], m["work"]])
        else:
            for sh in (names[i], infos[i], cards[i], pics[i]):
                sh._element.getparent().remove(sh._element)
    for ph in team.placeholders:
        if ph.placeholder_format.idx == 0:
            ph.text_frame.text = "КОМАНДА"

    # 10 — история, выбор задачи, сложности
    set_lines(by_text(story, "Расскажите, как вы собрались"), [TEAM["history"]])
    set_lines(by_text(story, "Что вас вдохновило"), [
        "Реальные данные города — 14 ГБ журналов за 8 лет — и польза, видная сразу: меньше ложных выездов, "
        "ТО по состоянию."
    ])
    set_lines(by_text(story, "Расскажите о самых интересных"), [
        "Данные без меток и почти без реальных пожаров: метки вывели из журналов, для пожара и НСД — правила "
        "с вероятностью. Каскады по сотне каналов — склейка в эпизоды.",
    ])
    for ph in story.placeholders:
        if ph.placeholder_format.idx == 0:
            ph.text_frame.text = "О КОМАНДЕ И ЗАДАЧЕ"

    # 11 — коротко о решении
    set_lines(by_text(brief, "Опишите в чем техническая"), [
        "Модульная платформа: независимые модули и сменные адаптеры",
        "No-code: датчики, зоны, регламенты — без программиста",
        "Поток СМВУ (Kafka) → TimescaleDB → LightGBM на 24 ч",
        "Склейка тревог в эпизоды ×42, причина и приоритет",
        "Django, React, Celery, Grafana; AD, 8 ролей, TLS",
    ])
    set_lines(by_text(brief, "Опишите ваши идеи"), [
        "Диспетчерам: меньше ложных выездов, очередь по опасности",
        "Эксплуатации: ТО по состоянию, графики в формах заказчика",
        "Внедрению: одна команда, закрытый контур",
        "Тиражированию: новый район и датчики без кода",
        "Эффект: ≈ 9,6 млн ₽ в год на район",
    ])


def io_bytes(blob):
    import io

    return io.BytesIO(blob)


# ---------------------------------------------------------------- описательная часть


def slide_problem(prs):
    s = new_slide(prs)
    header(s, "Проблема и данные")
    card(s, 1.0, 3.2, 10.2, 7.2, "Поток тревог", [
        "• 90 250 сигналов за 30 суток",
        "• Каскад — сотня тревог за секунду",
    ], size=14)
    card(s, 11.8, 3.2, 10.2, 7.2, "Ложные выезды", [
        "• Тревоги в основном ложные",
        "• Выезды «на всякий случай»",
    ], size=14)
    card(s, 22.6, 3.2, 10.2, 7.2, "ТО по календарю", [
        "• Без учёта состояния",
        "• Графики вручную в XLSX",
    ], size=14)
    rect(s, 1.0, 11.0, 31.8, 6.6, fill=DARK, radius=0.05)
    text(s, 1.6, 11.4, 30.6, 1.0, ["Данные заказчика"], size=14, bold=True, color=WHITE)
    stats = [("14 ГБ", "журналов СМВУ\nза 8 лет"), ("11 485", "каналов,\n19 типов датчиков"),
             ("95", "объектов\nрайона"), ("1 008", "датчиков метана\nв ППР"), ("24", "объекта\nв графике ТО")]
    for i, (big, small) in enumerate(stats):
        x = 1.6 + i * 6.2
        text(s, x, 12.6, 5.8, 1.6, [big], size=30, bold=True, color=PINK)
        text(s, x, 14.4, 5.8, 2.4, small.split("\n"), size=12, color=WHITE)


def slide_cycle(prs):
    s = new_slide(prs)
    header(s, "Решение: замкнутый цикл")
    steps = [
        ("Предсказать", "риск на 24 ч"),
        ("Объединить", "сотня тревог → 1 эпизод"),
        ("Расставить приоритет", "очередь по опасности"),
        ("Объяснить", "причина и факторы"),
        ("Предложить действие", "чек-лист, заявка"),
        ("Получить результат", "решение, отчёт бригады"),
        ("Улучшить модель", "обучение на решениях"),
    ]
    boxes = []
    for i, (title, desc) in enumerate(steps):
        x = 1.0 + i * 4.6
        b = rect(s, x, 4.2, 4.1, 2.2, fill=PINK if i in (0, 6) else WHITE, radius=0.15)
        label(b, [title], size=12, bold=True, color=WHITE if i in (0, 6) else PURPLE)
        boxes.append(b)
        rect(s, x, 7.0, 4.1, 3.4, fill=WHITE, radius=0.06)
        text(s, x + 0.25, 7.3, 3.6, 3.0, [desc], size=13, color=INK, align=PP_ALIGN.CENTER)
    for a, b in zip(boxes, boxes[1:]):
        arrow(s, a, b, elbow=False, color=WHITE, width=2)
    loop = rect(s, 1.0, 11.2, 31.8, 1.2, fill=None, line=PINK, radius=0.5, line_width=2.0)
    label(loop, ["↻ решение человека возвращается в систему как обратная связь"], size=12, color=WHITE)
    # главная демонстрация — цепочкой
    demo = [("47", "тревог на объекте"), ("1", "карточка"), ("3", "гипотезы причины"), ("1", "действие"), ("1", "заявка")]
    prev = None
    for i, (big, small) in enumerate(demo):
        x = 1.0 + i * 6.45
        b = rect(s, x, 13.4, 5.6, 4.2, fill=DARK, radius=0.08)
        label(b, [(big, {"size": 32, "bold": True, "color": PINK}), (small, {"size": 12, "color": WHITE})])
        if prev:
            arrow(s, prev, b, elbow=False, color=WHITE, width=2)
        prev = b


def slide_features(prs):
    s = new_slide(prs)
    header(s, "Функции")
    features = [
        ("Мониторинг", "карта зоны, 5 режимов"),
        ("Эпизоды", "сигналы → карточка ×42"),
        ("Прогноз", "отказ, метан, вода на 24 ч"),
        ("Пожар и НСД", "вероятность по истории"),
        ("Здание по снимку", "ИИ, планы этажей"),
        ("Заявки", "от черновика до отчёта"),
        ("Инженер ТО", "план по состоянию"),
        ("Графики ТО и ППР", "в формах заказчика"),
        ("Конструктор", "датчик без программиста"),
        ("Обучение моделей", "на решениях диспетчеров"),
        ("Аналитика", "смена, прогнозы, отчёты"),
        ("Учения", "полигон и симулятор"),
    ]
    for i, (title, desc) in enumerate(features):
        col, row = i % 4, i // 4
        x, y = 1.0 + col * 8.0, 3.2 + row * 4.9
        rect(s, x, y, 7.6, 4.5, fill=WHITE, radius=0.06)
        rect(s, x + 0.4, y + 0.55, 0.35, 0.35, fill=PINK, shape=MSO_SHAPE.OVAL)
        text(s, x + 0.95, y + 0.3, 6.4, 0.9, [title], size=15, bold=True, color=PURPLE)
        text(s, x + 0.4, y + 1.7, 6.8, 2.4, [desc], size=14, color=INK)


def slide_roles(prs):
    s = new_slide(prs)
    header(s, "8 ролей — 8 рабочих мест")
    roles = [
        ("Диспетчер ОДС", "весь район"),
        ("Диспетчер подразделения", "своя зона"),
        ("Руководитель", "эскалации, утверждение"),
        ("Аналитик", "модели и данные"),
        ("Инженер ТО", "план и графики ТО"),
        ("Ремонтная бригада", "заявки с телефона"),
        ("Наблюдатель", "только просмотр"),
        ("Администратор", "роли и интеграции"),
    ]
    for i, (role, what) in enumerate(roles):
        y = 3.2 + i * 1.8
        rect(s, 1.0, y, 11.6, 1.55, fill=WHITE, radius=0.12)
        text(s, 1.4, y + 0.12, 11.0, 0.7, [role], size=12, bold=True, color=PURPLE)
        text(s, 1.4, y + 0.78, 11.0, 0.7, [what], size=10.5, color=GREY)
    picture(s, IMG / "guides" / "engineer_1_workspace.png", 13.6, 3.3, w=18.8)
    # вертикаль — схемой
    chain = ["Бригада", "Диспетчерская объекта", "ОДС", "Руководство"]
    prev = None
    for i, name in enumerate(chain):
        b = rect(s, 13.6 + i * 4.85, 16.0, 4.3, 1.6, fill=PINK if i == 3 else WHITE, radius=0.3)
        label(b, [name], size=10.5, bold=True, color=WHITE if i == 3 else PURPLE)
        if prev:
            arrow(s, prev, b, elbow=False, color=WHITE, width=2)
        prev = b
    text(s, 13.6, 17.7, 19.0, 0.8, ["карточка без реакции в норматив поднимается выше"], size=10, color=WHITE)


def _lanes(s, x, y, w, lanes, lane_h, title_w=3.2):
    rect(s, x, y, w, lane_h * len(lanes), fill=WHITE, radius=0.02)
    for i, name in enumerate(lanes):
        ly = y + i * lane_h
        band = rect(s, x, ly, title_w, lane_h, fill=LANE if i % 2 == 0 else WHITE, line=LAVENDER, radius=0.0,
                    shape=MSO_SHAPE.RECTANGLE, line_width=0.5)
        label(band, [name], size=10, bold=True, color=PURPLE)
        if i:
            line = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Cm(x), Cm(ly), Cm(x + w), Cm(ly))
            line.line.color.rgb = LAVENDER
            line.line.width = Pt(0.5)


def _task(s, x, y, text_, w=3.4, h=1.5, fill=PINK_LIGHT):
    b = rect(s, x, y, w, h, fill=fill, line=PURPLE, radius=0.18, line_width=1.0)
    label(b, [text_], size=10, color=INK)
    return b


def _event(s, x, y, end=False, text_=None):
    e = rect(s, x, y, 0.9, 0.9, fill=WHITE, line=DARK if end else PURPLE, shape=MSO_SHAPE.OVAL, line_width=3.0 if end else 1.5)
    if text_:
        text(s, x - 1.0, y + 0.95, 2.9, 0.8, [text_], size=8, color=GREY, align=PP_ALIGN.CENTER)
    return e


def _gateway(s, x, y, text_=None):
    g = rect(s, x, y, 1.2, 1.2, fill=WHITE, line=PURPLE, shape=MSO_SHAPE.DIAMOND, line_width=1.5)
    label(g, ["×"], size=12, bold=True, color=PURPLE)
    if text_:
        text(s, x - 1.1, y - 0.75, 3.4, 0.8, [text_], size=8, color=GREY, align=PP_ALIGN.CENTER)
    return g


def slide_bpmn_incident(prs):
    s = new_slide(prs)
    header(s, "BPMN: от сигнала до решения")
    lanes = ["СМВУ и система", "Диспетчер", "Руководитель", "Бригада", "Аналитик"]
    X, Y, W, H = 1.0, 2.9, 31.8, 3.05
    _lanes(s, X, Y, W, lanes, H)
    row = lambda i, h=1.5: Y + i * H + (H - h) / 2  # noqa: E731
    start = _event(s, 4.6, row(0, 0.9), text_="сигнал датчика")
    norm = _task(s, 6.2, row(0), "Нормализация, эпизод")
    card_ = _task(s, 10.2, row(0), "Карточка: причина, приоритет")
    notify = _task(s, 14.2, row(0), "Уведомление зоне")
    ack = _task(s, 14.2, row(1), "Принять («кто первый»)")
    g1 = _gateway(s, 18.4, row(1, 1.2), "в норматив?")
    esc = _task(s, 18.0, row(2), "Эскалация, перехват")
    check = _task(s, 20.4, row(1), "Проверка")
    decide = _task(s, 24.4, row(1), "Решение")
    g2 = _gateway(s, 28.6, row(1, 1.2), "нужна работа?")
    draft = _task(s, 28.2, row(2), "Утвердить заявку")
    work = _task(s, 28.2, row(3), "Работа и отчёт")
    end = _event(s, 31.3, row(1, 0.9) + 0.0, end=True)
    label_ = _task(s, 20.4, row(4), "Метка обучения")
    retrain = _task(s, 24.4, row(4), "Переобучение")
    for a, b, sa, sb in [
        (start, norm, "right", "left"), (norm, card_, "right", "left"), (card_, notify, "right", "left"),
        (notify, ack, "bottom", "top"), (ack, g1, "right", "left"), (g1, check, "right", "left"),
        (g1, esc, "bottom", "top"), (esc, check, "right", "bottom"), (check, decide, "right", "left"),
        (decide, g2, "right", "left"), (g2, draft, "bottom", "top"), (draft, work, "bottom", "top"),
        (g2, end, "right", "left"), (decide, label_, "bottom", "top"), (label_, retrain, "right", "left"),
    ]:
        arrow(s, a, b, sa, sb, width=1.25)


def slide_bpmn_maintenance(prs):
    s = new_slide(prs)
    header(s, "BPMN: процесс ТО")
    lanes = ["Система", "Инженер ТО", "Руководитель", "Бригада"]
    X, Y, W, H = 1.0, 3.0, 31.8, 3.6
    _lanes(s, X, Y, W, lanes, H)
    row = lambda i, h=1.5: Y + i * H + (H - h) / 2  # noqa: E731
    start = _event(s, 4.6, row(0, 0.9), text_="начало года")
    reg = _task(s, 6.2, row(0), "Реестр и регламент")
    gen = _task(s, 10.2, row(0), "График ТО и ППР")
    check = _task(s, 10.2, row(1), "Сверка, правки")
    approve = _task(s, 14.2, row(2), "Утвердить график")
    due = _task(s, 18.2, row(0), "Работы месяца и рекомендации")
    plan = _task(s, 18.2, row(1), "В план → заявка")
    appr_wo = _task(s, 22.2, row(2), "Утвердить заявку")
    do = _task(s, 22.2, row(3), "Выполнить, отчёт")
    g = _gateway(s, 26.6, row(3, 1.2), "требует ремонта?")
    rec = _task(s, 26.2, row(0), "Рекомендация")
    reg2 = _task(s, 29.0, row(1), "Состояние в реестр", w=3.2)
    end = _event(s, 31.4, row(3, 0.9), end=True)
    for a, b, sa, sb in [
        (start, reg, "right", "left"), (reg, gen, "right", "left"), (gen, check, "bottom", "top"),
        (check, approve, "right", "left"), (approve, due, "right", "bottom"), (due, plan, "bottom", "top"),
        (plan, appr_wo, "right", "left"), (appr_wo, do, "bottom", "top"), (do, g, "right", "left"),
        (g, rec, "top", "bottom"), (g, end, "right", "left"), (do, reg2, "top", "bottom"), (rec, plan, "left", "right"),
    ]:
        arrow(s, a, b, sa, sb, width=1.25)


def slide_ml(prs):
    s = new_slide(prs)
    header(s, "ML-конвейер")
    top = [
        ("Журналы СМВУ", "2019–2026"),
        ("Нормализация", "профили датчиков"),
        ("Витрины", "сутки × канал"),
        ("Признаки", "окна 1–90 суток"),
        ("LightGBM", "обучение до 2025"),
    ]
    bottom = [
        ("Переобучение", "раз в неделю"),
        ("Метки из решений", "проверяет аналитик"),
        ("Журнал и исходы", "что сбылось"),
        ("Уровни риска", "пороги по точности"),
        ("Тест 2026", "отложенный год"),
    ]
    tb, bb = [], []
    for i, (t, d) in enumerate(top):
        b = rect(s, 1.0 + i * 6.4, 3.3, 5.6, 3.0, fill=WHITE, radius=0.08)
        label(b, [(t, {"bold": True, "color": PURPLE, "size": 14}), (d, {"size": 12})])
        tb.append(b)
    for i, (t, d) in enumerate(bottom):
        b = rect(s, 1.0 + i * 6.4, 8.1, 5.6, 3.0, fill=PINK_LIGHT if i == 0 else WHITE, radius=0.08)
        label(b, [(t, {"bold": True, "color": PURPLE, "size": 14}), (d, {"size": 12})])
        bb.append(b)
    for a, b in zip(tb, tb[1:]):
        arrow(s, a, b, elbow=False, color=WHITE, width=2)
    arrow(s, tb[-1], bb[-1], "bottom", "top", elbow=False, color=WHITE, width=2)
    for a, b in zip(bb[::-1], bb[::-1][1:]):
        arrow(s, a, b, "left", "right", elbow=False, color=WHITE, width=2)
    arrow(s, bb[0], tb[0], "top", "bottom", elbow=False, color=PINK, width=2, dash=True)
    card(s, 1.0, 12.0, 15.6, 5.8, "Принципы", [
        "• Прогноз начала события",
        "• Одни признаки в обучении и в работе",
        "• Сравнение с простым правилом",
    ], size=13)
    card(s, 17.2, 12.0, 15.6, 5.8, "Пожар и НСД — без модели", [
        "• Подтверждённых событий в данных нет",
        "• Правила + вероятность по архиву",
        "• Пожар: Brier 0,249 против 0,296",
    ], size=13)


def slide_results(prs):
    s = new_slide(prs)
    header(s, "Результаты: тест 2026 года")
    rows = [
        ["Задача", "ROC-AUC", "PR-AUC (случайная)", "Высокий: точность / полнота", "Правило: точность / полнота"],
        ["Отказ датчика, 24 ч", "0,90", "0,112 (0,006)", "0,21 / 0,09 · 32 в сутки", "0,10 / 0,45"],
        ["Метан ≥ 1 %, 24 ч", "0,94", "0,188 (0,004)", "0,08 / 0,74 · 20 в сутки", "0,07 / 0,39"],
        ["Подтопление, 24 ч", "0,82", "0,179 (0,040)", "0,15 / 0,46 · 21 в сутки", "0,13 / 0,54"],
    ]
    tbl = s.shapes.add_table(len(rows), len(rows[0]), Cm(1.0), Cm(3.3), Cm(31.8), Cm(5.2)).table
    widths = [7.0, 3.6, 5.6, 8.0, 7.6]
    for i, w in enumerate(widths):
        tbl.columns[i].width = Cm(w)
    for r, row in enumerate(rows):
        for c, val in enumerate(row):
            cell = tbl.cell(r, c)
            cell.text = val
            cell.fill.solid()
            cell.fill.fore_color.rgb = PURPLE if r == 0 else (WHITE if r % 2 else LANE)
            for p in cell.text_frame.paragraphs:
                for run in p.runs:
                    _font(run, 11, r == 0, WHITE if r == 0 else INK)
    stats = [
        ("×33", "точнее случайного выбора"),
        ("74 %", "превышений метана ловит модель (правило — 39 %)"),
        ("12 ч", "упреждение прогноза отказа"),
        ("×42", "сжатие потока тревог"),
    ]
    for i, (big, small) in enumerate(stats):
        x = 1.0 + i * 8.0
        rect(s, x, 9.3, 7.6, 4.2, fill=WHITE, radius=0.06)
        text(s, x + 0.4, 9.6, 6.8, 1.5, [big], size=28, bold=True, color=PINK)
        text(s, x + 0.4, 11.2, 6.8, 2.2, [small], size=12.5, color=INK)
    rect(s, 1.0, 14.2, 31.8, 3.0, fill=DARK, radius=0.06)
    text(s, 1.5, 14.5, 30.8, 2.5, [
        ("Уровни риска выбраны по точности", {"bold": True, "color": PINK, "size": 13}),
        ("«Критический» — немедленная реакция, «высокий» — плановая проверка", {"color": WHITE}),
    ], size=12)


def slide_fire(prs):
    s = new_slide(prs)
    header(s, "Пожар и НСД, камеры")
    picture(s, IMG / "fire_probability.png", 1.2, 3.4, w=16.6)
    card(s, 18.6, 3.2, 14.2, 7.0, "Вероятность по истории", [
        "• Пожар: 7 % → 42 % → 63 % по индексу",
        "• В карточке — число похожих случаев",
        "• НСД: главное — проверка",
    ], size=13)
    card(s, 18.6, 10.8, 14.2, 7.0, "Проверка по камерам", [
        "• Ближайшие камеры к тревоге",
        "• Кадр на момент тревоги и сейчас",
        "• Просмотр — в хронологии",
    ], size=13)


def slide_episode(prs):
    s = new_slide(prs)
    header(s, "Эпизод вместо сотни тревог")
    picture(s, IMG / "incident_cameras.png", 1.2, 3.4, w=18.4)
    card(s, 20.4, 3.2, 12.4, 14.6, "Что видит диспетчер", [
        "• Сколько сигналов склеено",
        "• Вероятные причины с весами",
        "• Приоритет 0–100",
        "• Чек-лист «что делать»",
        "• Заявка одной кнопкой",
    ], size=14)


def slide_floors(prs):
    s = new_slide(prs)
    header(s, "Здание по снимку и планы этажей")
    picture(s, IMG / "floor_georef.png", 1.2, 3.4, w=19.0)
    text(s, 1.2, 15.6, 19.0, 1.4, ["План этажа на снимке по 5 контрольным точкам"], size=12, color=WHITE)
    picture(s, IMG / "ai_segment.png", 21.0, 3.4, w=11.8)
    # как это работает — схемой
    steps = ["ИИ: контур здания", "План на снимок", "Этажи", "Датчики на этаж"]
    prev = None
    for i, t in enumerate(steps):
        b = rect(s, 21.0, 12.2 + i * 1.45, 11.8, 1.15, fill=PINK if i == 0 else WHITE, radius=0.3)
        label(b, [t], size=12, bold=True, color=WHITE if i == 0 else PURPLE)
        if prev:
            arrow(s, prev, b, "bottom", "top", elbow=False, color=WHITE, width=1.5)
        prev = b


def slide_accidents(prs):
    s = new_slide(prs)
    header(s, "Аварии и маршрут нарушителя")
    picture(s, IMG / "intrusion_route.png", 1.2, 3.4, h=14.4)
    picture(s, IMG / "monitoring_floors.png", 13.6, 3.4, w=11.4)
    card(s, 13.6, 12.0, 11.4, 5.8, "Маршрут нарушителя", [
        "• Люк → двери → движение, по шагам",
        "• Где сейчас и куда идёт",
    ], size=13)
    groups = ["Пожар", "Наводнение", "Газ", "Проникновение", "Температура"]
    rect(s, 25.6, 3.2, 7.2, 14.6, fill=WHITE, radius=0.06)
    text(s, 26.0, 3.55, 6.4, 1.0, ["5 групп аварий"], size=14, bold=True, color=PURPLE)
    for i, g in enumerate(groups):
        b = rect(s, 26.0, 5.0 + i * 2.1, 6.4, 1.7, fill=PINK_LIGHT if i < 4 else PINK, radius=0.3)
        label(b, [g], size=12.5, bold=True, color=WHITE if i == 4 else PURPLE)
    text(s, 26.0, 15.6, 6.4, 2.0, ["Потеря связи или питания — инцидент"], size=10.5, color=GREY)


def slide_maintenance(prs):
    s = new_slide(prs)
    header(s, "Инженер ТО: графики")
    picture(s, IMG / "schedules.png", 1.2, 3.4, w=19.0)
    stats = [("91 %", "совпадение периодичности с графиком заказчика"), ("47 / 62", "пик работ в месяц: система / заказчик")]
    for i, (big, small) in enumerate(stats):
        y = 3.2 + i * 4.3
        rect(s, 21.0, y, 11.8, 3.9, fill=WHITE, radius=0.06)
        text(s, 21.4, y + 0.3, 11.0, 1.5, [big], size=28, bold=True, color=PINK)
        text(s, 21.4, y + 1.9, 11.0, 1.8, [small], size=12.5, color=INK)
    steps = ["Регламент", "График", "Утверждение", "Заявки"]
    prev = None
    for i, t in enumerate(steps):
        b = rect(s, 21.0, 12.0 + i * 1.5, 11.8, 1.2, fill=PINK if i == 0 else WHITE, radius=0.3)
        label(b, [t], size=12, bold=True, color=WHITE if i == 0 else PURPLE)
        if prev:
            arrow(s, prev, b, "bottom", "top", elbow=False, color=WHITE, width=1.5)
        prev = b


def slide_constructor(prs):
    s = new_slide(prs)
    header(s, "Конструктор датчиков: no-code")
    layers = [("Сообщение", "JSON, CSV, строка"), ("Шаблон формата", "без кода"),
              ("Единое событие", "для всех источников"), ("Профиль датчика", "пороги, коды"),
              ("Конвейер", "прогноз, карточки")]
    prev = None
    for i, (t, d) in enumerate(layers):
        b = rect(s, 1.0 + i * 6.4, 3.3, 5.6, 2.8, fill=PINK if i == 1 else WHITE, radius=0.1)
        label(b, [(t, {"bold": True, "size": 13, "color": WHITE if i == 1 else PURPLE}), (d, {"size": 11, "color": WHITE if i == 1 else INK})])
        if prev:
            arrow(s, prev, b, elbow=False, color=WHITE, width=2)
        prev = b
    picture(s, IMG / "constructor.png", 1.2, 7.1, w=14.6)
    rows = [["Что подключить", "Код"], ["Датчик нового формата", "не нужен"], ["Новый тип датчика", "не нужен"],
            ["Регламент ТО", "не нужен"], ["Система заказчика", "не нужен"], ["Новая задача прогноза", "немного"]]
    tbl = s.shapes.add_table(len(rows), 2, Cm(16.8), Cm(7.0), Cm(16.0), Cm(8.0)).table
    for i, w in enumerate((11.0, 5.0)):
        tbl.columns[i].width = Cm(w)
    for r, row in enumerate(rows):
        for c, val in enumerate(row):
            cell = tbl.cell(r, c)
            cell.text = val
            cell.fill.solid()
            cell.fill.fore_color.rgb = PURPLE if r == 0 else (WHITE if r % 2 else LANE)
            for p in cell.text_frame.paragraphs:
                for run in p.runs:
                    _font(run, 12, r == 0, WHITE if r == 0 else (PINK if c == 1 and val == "не нужен" else INK))


def slide_architecture(prs):
    s = new_slide(prs)
    header(s, "Модульная архитектура")
    # источники → адаптеры → модули ядра → клиенты, одной схемой
    sources = ["СМВУ (Kafka)", "Файлы CSV/XLSX", "Новые датчики", "Погода"]
    clients = ["Веб и телефон", "REST API", "Help desk, VMS, AD", "Grafana"]
    src_boxes, cl_boxes = [], []
    for i, t in enumerate(sources):
        b = rect(s, 1.0, 3.4 + i * 2.5, 5.6, 2.0, fill=WHITE, radius=0.2)
        label(b, [t], size=12, bold=True, color=PURPLE)
        src_boxes.append(b)
    for i, t in enumerate(clients):
        b = rect(s, 27.2, 3.4 + i * 2.5, 5.6, 2.0, fill=WHITE, radius=0.2)
        label(b, [t], size=12, bold=True, color=PURPLE)
        cl_boxes.append(b)
    adapters = rect(s, 7.2, 3.4, 3.9, 9.5, fill=PINK, radius=0.15)
    label(adapters, ["Адаптеры", ("сменные", {"size": 10, "color": WHITE})], size=12, bold=True, color=WHITE)
    core = rect(s, 11.6, 3.4, 14.6, 9.5, fill=DARK, radius=0.05)
    text(s, 12.0, 3.6, 13.8, 0.9, ["Ядро: независимые модули"], size=13, bold=True, color=WHITE)
    modules = ["Приём", "Нормализация", "Прогноз", "Инциденты", "Заявки", "ТО", "Роли и зоны", "Аудит", "Обучение"]
    for i, m in enumerate(modules):
        col, row = i % 3, i // 3
        b = rect(s, 12.0 + col * 4.65, 4.8 + row * 2.6, 4.3, 2.1, fill=WHITE, radius=0.15)
        label(b, [m], size=12, bold=True, color=PURPLE)
    for b in src_boxes:
        arrow(s, b, adapters, elbow=False, color=WHITE, width=1.5)
    arrow(s, adapters, core, elbow=False, color=WHITE, width=2)
    for b in cl_boxes:
        arrow(s, core, b, elbow=False, color=WHITE, width=1.5)
    # no-code — полосой под схемой
    nocode = [("Конструктор", "датчики и форматы"), ("Интерфейс", "зоны, объекты, регламенты"),
              (".env", "системы заказчика"), ("Полигон", "сценарии учений")]
    text(s, 1.0, 13.5, 31.8, 0.9, ["No-code: что меняется без программиста"], size=13, bold=True, color=WHITE)
    for i, (t, d) in enumerate(nocode):
        b = rect(s, 1.0 + i * 8.0, 14.6, 7.6, 3.0, fill=WHITE, radius=0.08)
        label(b, [(t, {"bold": True, "size": 13, "color": PINK}), (d, {"size": 11.5})])


def slide_quality(prs):
    s = new_slide(prs)
    header(s, "Надёжность и безопасность")
    stats = [("47 мс", "ответ при 20 пользователях"), ("≈5 с", "прогноз по всем объектам (ТЗ ≤ 300 с)"),
             ("242", "автотеста"), ("≤ 4 ч", "восстановление из копии")]
    for i, (big, small) in enumerate(stats):
        x = 1.0 + i * 8.0
        rect(s, x, 3.3, 7.6, 4.2, fill=WHITE, radius=0.06)
        text(s, x + 0.4, 3.6, 6.8, 1.5, [big], size=28, bold=True, color=PINK)
        text(s, x + 0.4, 5.2, 6.8, 2.2, [small], size=12.5, color=INK)
    card(s, 1.0, 8.2, 15.6, 9.6, "Безопасность (ТЗ §11)", [
        "• TLS, вход через AD",
        "• 8 ролей и зона на каждом запросе",
        "• Журнал действий и просмотров",
        "• Чек-лист ввода в эксплуатацию",
    ], size=14)
    card(s, 17.2, 8.2, 15.6, 9.6, "Эксплуатация", [
        "• Grafana и оповещения Prometheus",
        "• Ежедневные резервные копии",
        "• Модели готовы с первого дня",
        "• 120 страниц документации",
    ], size=14)


def slide_map(prs):
    s = new_slide(prs)
    header(s, "Карта мониторинга")
    picture(s, IMG / "monitoring_district.png", 1.2, 3.4, w=19.0)
    card(s, 21.0, 3.2, 11.8, 8.4, "Главный экран каждой роли", [
        "• Своя зона в цвете",
        "• Карточки и заявки на объектах",
        "• 5 режимов окраски",
    ], size=14)
    card(s, 21.0, 12.2, 11.8, 5.6, "Подложки и этажи", [
        "• Карта или спутник",
        "• Планы этажей объекта",
    ], size=14)


def slide_mobile(prs):
    s = new_slide(prs)
    header(s, "Мобильное приложение в браузере")
    picture(s, IMG / "mobile_monitoring.png", 1.2, 3.4, h=14.2)
    picture(s, IMG / "guides" / "brigade_4_complete.png", 15.3, 3.4, h=14.2)
    card(s, 22.8, 3.2, 10.0, 8.2, "Без магазина приложений", [
        "• Ссылка → «На главный экран»",
        "• Тот же вход и роли",
    ], size=14)
    card(s, 22.8, 12.0, 10.0, 5.8, "На объекте", [
        "• Бригада: заявки и отчёт",
        "• Руководитель: утверждение",
    ], size=14)


def slide_training(prs):
    s = new_slide(prs)
    header(s, "Учебный контур и учения")
    picture(s, IMG / "exercise.png", 1.2, 3.4, w=19.0)
    steps = [("Руководитель", "выбирает сценарий"), ("Симулятор", "11 сценариев, 6 готовых"),
             ("Полигон", "настоящие карточки"), ("Смена", "реагирует как в жизни"), ("Разбор", "норматив, решения")]
    prev = None
    for i, (t, d) in enumerate(steps):
        b = rect(s, 21.0, 3.2 + i * 2.95, 11.8, 2.4, fill=PINK if i == 0 else WHITE, radius=0.15)
        label(b, [(t, {"bold": True, "size": 13, "color": WHITE if i == 0 else PURPLE}),
                  (d, {"size": 11.5, "color": WHITE if i == 0 else INK})])
        if prev:
            arrow(s, prev, b, "bottom", "top", elbow=False, color=WHITE, width=1.5)
        prev = b


def slide_tz(prs):
    s = new_slide(prs)
    header(s, "Соответствие ТЗ")
    rows = [["Раздел", "Требование", "Реализация"],
            ["§3, §8", "Планирование ТО, черновики заявок", "План и графики ТО и ППР, заявки"],
            ["§4, §6", "Отказы, пожар, подтопление, НСД", "3 модели на 24 ч, индикаторы, маршрут нарушителя"],
            ["§5", "Не заменяет диспетчера", "Решение — за человеком"],
            ["§7", "CSV, XLSX, JSON, XML, GeoJSON, WKT", "Импорт, REST, GeoJSON/WKT, планы этажей"],
            ["§9, §11", "≤ 300 с, 20 пользователей", "≈5 с прогноз, 47 мс ответ"],
            ["§10", "Дашборд, карта, журнал, уведомления", "Мониторинг, журнал прогнозов, WebSocket"],
            ["§11", "TLS, LDAP/AD, RBAC, журнал", "TLS, AD, 8 ролей и зоны, журнал действий"],
            ["§12", "Пути «пожар» и «подтопление»", "Карточка → решение → заявка → обучение"],
            ["§13", "Метео, только чтение, контуры", "Open-Meteo, адаптеры на чтение, контуры данных"],
            ["§14", "Документация, методы, библиотеки", "120 страниц, руководства по ролям, библиотеки"]]
    tbl = s.shapes.add_table(len(rows), 3, Cm(1.0), Cm(3.2), Cm(31.8), Cm(14.6)).table
    for i, w in enumerate((3.2, 13.0, 15.6)):
        tbl.columns[i].width = Cm(w)
    for r, row in enumerate(rows):
        for c, val in enumerate(row):
            cell = tbl.cell(r, c)
            cell.text = val
            cell.fill.solid()
            cell.fill.fore_color.rgb = PURPLE if r == 0 else (WHITE if r % 2 else LANE)
            for p in cell.text_frame.paragraphs:
                for run in p.runs:
                    _font(run, 11.5, r == 0 or c == 0, WHITE if r == 0 else INK)


def slide_business(prs):
    s = new_slide(prs)
    header(s, "Эффект для эксплуатации")
    items = [
        ("Меньше ложных выездов", "причина видна до выезда"),
        ("Раньше о рисках", "прогноз на сутки"),
        ("ТО по состоянию", "ровная загрузка бригад"),
        ("Прозрачная смена", "нормативы и показатели"),
        ("Быстрое внедрение", "одна команда развёртывания"),
        ("Масштабирование", "новый район без кода"),
    ]
    for i, (t, d) in enumerate(items):
        col, row = i % 3, i // 3
        x, y = 1.0 + col * 10.7, 3.3 + row * 7.3
        rect(s, x, y, 10.2, 6.8, fill=WHITE, radius=0.06)
        text(s, x + 0.5, y + 0.4, 9.2, 1.0, [f"0{i + 1}"], size=22, bold=True, color=PINK)
        text(s, x + 0.5, y + 2.0, 9.2, 1.2, [t], size=16, bold=True, color=PURPLE)
        text(s, x + 0.5, y + 3.5, 9.2, 2.8, [d], size=14, color=INK)


# Экономический эффект: объёмы — из данных заказчика, ставки и доли — допущения (docs/delivery/economics.md)
ECON = {
    "signals_30d": 90_250,
    "cards_30d": 2_138,
    "trip_share": 0.05,  # доля карточек, по которым без системы выезжает бригада
    "trip_avoided": 0.30,  # доля этих выездов, которые снимает система
    "trip_cost": 15_000,  # ₽ за выезд бригады
    "signal_seconds": 10,  # просмотр сигнала без системы
    "card_minutes": 2,  # разбор карточки эпизода
    "hour_cost": 800,  # ₽ за час диспетчера с взносами
    "failures_per_day": 32 * 0.21,  # верные прогнозы отказа уровня «высокий» в сутки
    "replace_share": 0.10,  # доля отказов, требующих выезда на замену
    "replace_saving": 8_000,  # ₽: плановая замена вместо аварийной
}


def economics() -> list[tuple[str, float, str]]:
    e = ECON
    year = 365 / 30
    signals, cards = e["signals_30d"] * year, e["cards_30d"] * year
    trips = cards * e["trip_share"] * e["trip_avoided"]
    hours = signals * e["signal_seconds"] / 3600 - cards * e["card_minutes"] / 60
    replaced = e["failures_per_day"] * 365 * e["replace_share"]
    return [
        ("Ложные выезды", trips * e["trip_cost"], f"−{trips:.0f} выездов в год"),
        ("Время смены", hours * e["hour_cost"], f"−{hours:,.0f} ч в год".replace(",", " ")),
        ("Плановая замена", replaced * e["replace_saving"], f"{replaced:.0f} замен вместо аварийных"),
    ]


def _mln(value: float) -> str:
    return f"{value / 1e6:.1f}".replace(".", ",")


def slide_economics(prs):
    s = new_slide(prs)
    header(s, "Экономический эффект")
    parts = economics()
    total = sum(v for _, v, _ in parts)
    top = max(v for _, v, _ in parts)
    # столбцы — простая схема вклада
    for i, (name, value, note) in enumerate(parts):
        x = 1.0 + i * 7.2
        h = 9.0 * value / top
        rect(s, x, 3.4, 6.6, 12.2, fill=WHITE, radius=0.05)
        text(s, x + 0.3, 3.7, 6.0, 1.4, [f"{_mln(value)} млн ₽"], size=22, bold=True, color=PINK, align=PP_ALIGN.CENTER)
        bar = rect(s, x + 1.8, 14.0 - h, 3.0, h, fill=PINK if i == 0 else LAVENDER, radius=0.05)
        bar.shadow.inherit = False
        text(s, x + 0.3, 14.2, 6.0, 0.8, [name], size=13, bold=True, color=PURPLE, align=PP_ALIGN.CENTER)
        text(s, x + 0.3, 14.9, 6.0, 0.7, [note], size=10.5, color=GREY, align=PP_ALIGN.CENTER)
    rect(s, 23.0, 3.4, 9.8, 12.2, fill=DARK, radius=0.05)
    text(s, 23.4, 4.0, 9.0, 1.0, ["Итого на один район"], size=14, bold=True, color=WHITE)
    text(s, 23.4, 5.4, 9.0, 2.2, [f"≈ {_mln(total)} млн ₽"], size=36, bold=True, color=PINK)
    text(s, 23.4, 7.9, 9.0, 0.8, ["в год, 95 объектов"], size=13, color=WHITE)
    text(s, 23.4, 9.6, 9.0, 5.6, [
        ("Не вошло в расчёт:", {"bold": True, "color": PINK}),
        "• ущерб от предотвращённых аварий",
        "• ровная загрузка бригад по ТО",
        "• каждый новый район — ×N",
    ], size=12, color=WHITE)
    text(s, 1.0, 16.2, 31.8, 1.6, [
        "Объёмы — из журналов заказчика (1,1 млн сигналов и 26 тыс. карточек в год); ставки и доли — допущения: "
        "выезд 15 тыс. ₽, 5 % карточек с выездом, система снимает 30 %, час диспетчера 800 ₽. Расчёт — economics.md.",
    ], size=10, color=WHITE)


def slide_plans(prs):
    s = new_slide(prs)
    header(s, "Планы развития")
    stages = [("Пилот", ["AD и help desk", "поток СМВУ", "реестр оборудования"]),
              ("Данные", ["режим охраны для НСД", "журналы ремонтов", "уровень воды"]),
              ("Модели", ["прогноз на 7 суток", "разбор ошибок прогноза", "рекомендации «что если»"]),
              ("Продукт", ["офлайн для бригады", "учения на нескольких объектах", "другие районы"])]
    prev = None
    for i, (t, items) in enumerate(stages):
        x = 1.0 + i * 8.0
        b = rect(s, x, 3.4, 7.4, 1.8, fill=PINK if i == 0 else WHITE, radius=0.3)
        label(b, [t], size=14, bold=True, color=WHITE if i == 0 else PURPLE)
        if prev:
            arrow(s, prev, b, elbow=False, color=WHITE, width=2)
        prev = b
        rect(s, x, 5.8, 7.4, 8.4, fill=WHITE, radius=0.06)
        text(s, x + 0.4, 6.2, 6.6, 7.8, [f"• {it}" for it in items], size=14, color=INK)
    text(s, 1.0, 15.0, 31.8, 1.4, ["Модульное ядро: развитие без переделки"], size=14, bold=True, color=WHITE)


def slide_final(prs):
    s = new_slide(prs, "Титульный слайд")
    text(s, 1.2, 6.0, 20.0, 2.5, ["Спасибо!"], size=44, bold=True, color=WHITE)
    text(s, 1.2, 9.0, 20.0, 5.0, [
        TEAM["name"],
        "Репозиторий: [ссылка]",
        "Прототип: [ссылка]",
        "Документация и руководства: [ссылка]",
    ], size=16, color=WHITE)


def build() -> Path:
    prs = Presentation(str(TEMPLATE))
    template_slides = list(prs.slides)
    fill_mandatory(prs)
    for fn in (slide_problem, slide_cycle, slide_features, slide_roles, slide_bpmn_incident, slide_bpmn_maintenance,
               slide_ml, slide_results, slide_map, slide_floors, slide_episode, slide_fire, slide_accidents, slide_maintenance, slide_constructor,
               slide_architecture, slide_quality, slide_mobile, slide_training, slide_tz, slide_business, slide_economics,
               slide_plans, slide_final):
        fn(prs)
    # из шаблона остаются только обязательные 7–11, служебные и образцы удаляются
    for i, slide in enumerate(template_slides, start=1):
        if not 7 <= i <= 11:
            delete_slide(prs, slide)
    OUT.parent.mkdir(exist_ok=True)
    prs.save(str(OUT))
    return OUT


def to_pdf(path: Path) -> Path | None:
    try:
        import win32com.client
    except ImportError:
        return None
    pdf = path.with_suffix(".pdf")
    app = win32com.client.Dispatch("PowerPoint.Application")
    try:
        pres = app.Presentations.Open(str(path), WithWindow=False)
        pres.SaveAs(str(pdf), 32)
        pres.Close()
    finally:
        app.Quit()
    return pdf


if __name__ == "__main__":
    out = build()
    print(f"PPTX: {out}")
    pdf = to_pdf(out)
    if pdf:
        print(f"PDF:  {pdf}")
