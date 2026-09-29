"""
Сборка сопроводительной документации (ТЗ §14, §19): Markdown → DOCX → PDF.

Источник — docs/delivery/documentation.md; главы-методики подключаются директивой
`<!-- include: путь [strip-title] [shift=N] -->`, поэтому текст методик живёт в одном месте (docs/*.md)
и не расходится с документом. PDF делает Microsoft Word (COM); без Word собирается только DOCX.

    uv run --no-project --with python-docx --with pywin32 python docs/delivery/build.py
"""

from __future__ import annotations

import re
import sys
from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

HERE = Path(__file__).resolve().parent
OUT = HERE / "out"
TITLE = "Сервис прогнозирования инцидентов инженерных коллекторов"
# (источник, файл, подзаголовок титула): сопроводительная документация и отдельная инструкция для
# администраторов заказчика — её текст тот же, что глава в документации (docs/deployment.md)
DOCUMENTS = [
    (HERE / "documentation.md", "Сопроводительная документация.docx", "Сопроводительная документация"),
    (
        HERE.parent / "deployment.md",
        "Инструкция по развёртыванию.docx",
        "Инструкция по запуску и развёртыванию на серверах заказчика",
    ),
]
# Руководства пользователя по ролям: docs/guides/<роль>.md → out/Руководства/<Руководство — роль>.docx
GUIDES = [
    ("dispatcher-ods.md", "Диспетчер ОДС"),
    ("dispatcher-unit.md", "Диспетчер подразделения"),
    ("head.md", "Руководитель подразделения"),
    ("analyst.md", "Аналитик"),
    ("maintenance-engineer.md", "Инженер ТО"),
    ("brigade.md", "Ремонтная бригада"),
    ("observer.md", "Наблюдатель"),
    ("admin.md", "Администратор"),
]
DOCUMENTS += [
    (
        HERE.parent / "guides" / source,
        f"Руководства/Руководство пользователя — {role}.docx",
        f"Руководство пользователя: {role}",
    )
    for source, role in GUIDES
]
# Общее руководство для всех ролей и памятка для проверяющего (адреса и тестовые учётные записи)
DOCUMENTS += [
    (HERE.parent / "guides" / "general.md", "Общее руководство пользователя.docx", "Общее руководство пользователя"),
    (HERE.parent / "guides" / "reviewer.md", "Памятка для проверяющего.docx", "Памятка для проверяющего"),
]
ACCENT = RGBColor(0x1C, 0x4E, 0x9A)
INCLUDE = re.compile(r"<!--\s*include:\s*(\S+)(.*?)-->")


# ---------------- Markdown ----------------


def expand(path: Path) -> list[str]:
    lines = []
    for line in path.read_text(encoding="utf-8").splitlines():
        m = INCLUDE.fullmatch(line.strip())
        if not m:
            lines.append(line)
            continue
        target = (path.parent / m.group(1)).resolve()
        opts = m.group(2)
        shift = int(re.search(r"shift=(\d+)", opts).group(1)) if "shift=" in opts else 0
        body = expand(target)
        if "strip-title" in opts and body and body[0].startswith("# "):
            body = body[1:]
        base = target.parent
        out = []
        in_code = False
        for b in body:
            if b.startswith("```"):
                in_code = not in_code
            if not in_code and re.match(r"#+ ", b):
                b = "#" * shift + b
            # картинки включаемого файла — относительно его каталога
            b = re.sub(
                r"!\[([^\]]*)\]\((?!https?:)([^)]+)\)",
                lambda mm, base=base: f"![{mm.group(1)}]({(base / mm.group(2)).as_posix()})",
                b,
            )
            out.append(b)
        lines += out
    return lines


INLINE = re.compile(r"(\*\*[^*]+\*\*|`[^`]+`|\[[^\]]+\]\([^)]+\)|\*[^*\s][^*]*\*)")


def add_inline(paragraph, text: str, size: float | None = None, bold: bool = False) -> None:
    for part in INLINE.split(text):
        if not part:
            continue
        run_bold, italic, mono = bold, False, False
        if part.startswith("**") and part.endswith("**"):
            part, run_bold = part[2:-2], True
        elif part.startswith("`") and part.endswith("`"):
            part, mono = part[1:-1], True
        elif part.startswith("[") and "](" in part:
            part = part[1 : part.index("](")]
        elif part.startswith("*") and part.endswith("*") and len(part) > 2:
            part, italic = part[1:-1], True
        run = paragraph.add_run(part)
        run.bold = run_bold
        run.italic = italic
        if mono:
            run.font.name = "Consolas"
            run.font.size = Pt((size or 11) - 1)
        elif size:
            run.font.size = Pt(size)


def set_cell_shading(cell, hex_color: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color)
    tc_pr.append(shd)


def repeat_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    el = OxmlElement("w:tblHeader")
    el.set(qn("w:val"), "true")
    tr_pr.append(el)


def add_table(doc, rows: list[list[str]]) -> None:
    header, body = rows[0], rows[1:]
    table = doc.add_table(rows=1, cols=len(header))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    size = 9 if len(header) <= 4 else 8
    for i, text in enumerate(header):
        cell = table.rows[0].cells[i]
        cell.text = ""
        add_inline(cell.paragraphs[0], text, size=size, bold=True)
        set_cell_shading(cell, "E8EDF6")
    repeat_header(table.rows[0])
    # шапка не остаётся одна внизу страницы
    for cell in table.rows[0].cells:
        cell.paragraphs[0].paragraph_format.keep_with_next = True
    for row in body:
        cells = table.add_row().cells
        for i in range(len(header)):
            cells[i].text = ""
            add_inline(cells[i].paragraphs[0], row[i] if i < len(row) else "", size=size)
    doc.add_paragraph()


def add_code(doc, code: list[str]) -> None:
    table = doc.add_table(rows=1, cols=1)
    # блок кода не разрывается между страницами
    table.rows[0]._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
    cell = table.rows[0].cells[0]
    set_cell_shading(cell, "F3F4F7")
    p = cell.paragraphs[0]
    for i, line in enumerate(code):
        run = p.add_run(line)
        run.font.name = "Consolas"
        run.font.size = Pt(8.5)
        if i < len(code) - 1:
            run.add_break()
    doc.add_paragraph()


def add_image(doc, path: Path, caption: str) -> None:
    if not path.exists():
        print(f"нет изображения {path}", file=sys.stderr)
        return
    from docx.image.image import Image

    image = Image.from_file(str(path))
    # высокие снимки (телефон) — по высоте, иначе вылезут за страницу
    if image.px_height > image.px_width * 1.2:
        doc.add_picture(str(path), height=Cm(15))
    else:
        doc.add_picture(str(path), width=Cm(16.5))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap = doc.add_paragraph(caption)
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.runs[0].italic = True
    cap.runs[0].font.size = Pt(9)


def restart_numbering(doc, paragraph) -> None:
    """Новый нумерованный список начинается с 1: своя нумерация поверх абстрактной из стиля."""
    numbering = doc.part.numbering_part.element
    style_num = paragraph.style.element.pPr.numPr.numId.val
    abstract = numbering.num_having_numId(style_num).abstractNumId.val
    num = numbering.add_num(abstract)
    override = num.add_lvlOverride(ilvl=0)
    override.add_startOverride(1)
    num_pr = paragraph._p.get_or_add_pPr().get_or_add_numPr()
    num_pr.get_or_add_ilvl().val = 0
    num_pr.get_or_add_numId().val = num.numId


def render(doc, lines: list[str], base: Path) -> None:
    i = 0
    paragraph: list[str] = []
    in_ordered = False  # идёт ли нумерованный список верхнего уровня

    def flush():
        if paragraph:
            p = doc.add_paragraph()
            add_inline(p, " ".join(s.strip() for s in paragraph))
            paragraph.clear()

    first_chapter = True
    while i < len(lines):
        line = lines[i]
        stripped = line.strip()
        if stripped and not re.match(r"\s*([-*]|\d+\.) ", line):
            in_ordered = False  # список прервался — следующий нумерованный начнётся с 1
        if stripped.startswith("```"):
            flush()
            code = []
            i += 1
            while i < len(lines) and not lines[i].strip().startswith("```"):
                code.append(lines[i])
                i += 1
            add_code(doc, code)
        elif m := re.match(r"(#+) (.*)", line):
            flush()
            level = min(len(m.group(1)), 4)
            if level == 1 and not first_chapter:
                doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
            first_chapter = first_chapter and level != 1
            doc.add_heading(m.group(2).strip(), level=level)
        elif stripped.startswith("|"):
            flush()
            rows = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                cells = [c.strip() for c in lines[i].strip().strip("|").split("|")]
                if not all(re.fullmatch(r":?-+:?", c) for c in cells if c):
                    rows.append(cells)
                i += 1
            add_table(doc, rows)
            continue
        elif m := re.match(r"!\[([^\]]*)\]\(([^)]+)\)", stripped):
            flush()
            target = Path(m.group(2))
            add_image(doc, target if target.is_absolute() else base / target, m.group(1))
        elif m := re.match(r"(\s*)([-*]|\d+\.) (.*)", line):
            flush()
            depth = len(m.group(1)) // 2
            ordered = m.group(2)[0].isdigit()
            text = m.group(3)
            # продолжение пункта на следующих строках с отступом
            while (
                i + 1 < len(lines)
                and lines[i + 1].startswith("  " * (depth + 1))
                and not re.match(r"\s*([-*]|\d+\.) ", lines[i + 1])
            ):
                i += 1
                text += " " + lines[i].strip()
            style = "List Number" if ordered else "List Bullet"
            if depth:
                style += f" {min(depth + 1, 3)}"
            p = doc.add_paragraph(style=style)
            if ordered and not depth and not in_ordered:
                restart_numbering(doc, p)
            if not depth:
                in_ordered = ordered
            add_inline(p, text)
        elif not stripped or stripped.startswith("<!--"):
            flush()
        else:
            paragraph.append(line)
        i += 1
    flush()


# ---------------- Документ ----------------


def add_field(paragraph, instruction: str) -> None:
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = instruction
    sep = OxmlElement("w:fldChar")
    sep.set(qn("w:fldCharType"), "separate")
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.append(begin)
    run._r.append(instr)
    run._r.append(sep)
    run._r.append(end)


def setup_styles(doc) -> None:
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.1
    for level, size in ((1, 18), (2, 14), (3, 12), (4, 11)):
        style = doc.styles[f"Heading {level}"]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.color.rgb = ACCENT
        style.font.bold = True
        style.paragraph_format.space_before = Pt(14 if level <= 2 else 10)
        style.paragraph_format.space_after = Pt(6)
        style.paragraph_format.keep_with_next = True


def build(source: Path, docx: Path, subtitle: str) -> Path:
    docx.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    section = doc.sections[0]
    section.orientation = WD_ORIENT.PORTRAIT
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.left_margin = section.right_margin = Cm(2)
    section.top_margin = section.bottom_margin = Cm(2)
    setup_styles(doc)

    # титульный лист
    for _ in range(6):
        doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(TITLE)
    run.font.size, run.bold, run.font.color.rgb = Pt(24), True, ACCENT
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(subtitle)
    run.font.size = Pt(16)
    for _ in range(2):
        doc.add_paragraph()
    for text in (
        "Задача «8. ДЖКХ» — АО «Москоллектор»",
        "Прогнозирование отказов датчиков и раннее выявление рисков инцидентов:",
        "пожар, загазованность, подтопление, несанкционированный доступ",
    ):
        p = doc.add_paragraph(text)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for _ in range(10):
        doc.add_paragraph()
    p = doc.add_paragraph(f"Москва, {date.today():%Y}")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    # оглавление (обновляется Word при открытии и при экспорте в PDF)
    # заголовок оглавления — не стиль «Заголовок», иначе он попадёт в само оглавление
    p = doc.add_paragraph()
    run = p.add_run("Содержание")
    run.font.size, run.bold, run.font.color.rgb = Pt(18), True, ACCENT
    add_field(doc.add_paragraph(), 'TOC \\o "1-2" \\h \\z \\u')
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    render(doc, expand(source), source.parent)

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_field(footer, "PAGE")
    section.different_first_page_header_footer = True
    settings = doc.settings.element
    update = OxmlElement("w:updateFields")
    update.set(qn("w:val"), "true")
    settings.append(update)
    doc.save(docx)
    return docx


def to_pdf(docx: Path) -> Path | None:
    """Word обновляет оглавление и номера страниц, затем сохраняет PDF."""
    try:
        import win32com.client
    except ImportError:
        print("pywin32 не установлен — PDF не собран", file=sys.stderr)
        return None
    word = win32com.client.DispatchEx("Word.Application")
    word.Visible = False
    word.DisplayAlerts = 0
    try:
        document = word.Documents.Open(str(docx), ReadOnly=False)
        document.TablesOfContents(1).Update()
        document.Fields.Update()
        document.Save()
        pdf = docx.with_suffix(".pdf")
        document.SaveAs2(str(pdf), FileFormat=17)
        document.Close(False)
        return pdf
    finally:
        word.Quit()


if __name__ == "__main__":
    for source, name, subtitle in DOCUMENTS:
        path = build(source, OUT / name, subtitle)
        print(f"DOCX: {path}")
        pdf = to_pdf(path)
        if pdf:
            print(f"PDF:  {pdf}")
